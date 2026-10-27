import io
import uuid
from datetime import date, timedelta

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import Field
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from app.api.deps import Ctx, require
from app.core.db import get_db
from app.core.kpi_definitions import KPI_DEFINITIONS
from app.models.system import KPISnapshot, MaintenanceBudget, ProductionRecord
from app.models.tenancy import Department
from app.schemas.common import In
from app.services import kpi as svc
from app.services.common import audit, plant_today, utcnow

router = APIRouter(prefix="/kpis", tags=["kpi"])


def _scope(
    ctx: Ctx, db: Session, start: date | None, end: date | None, department_id: uuid.UUID | None, asset_id: uuid.UUID | None
) -> svc.Scope:
    today = plant_today(ctx.plant)
    end = end or today
    start = start or end.replace(day=1)
    if end < start or (end - start).days > 731:
        raise HTTPException(422, "Invalid period (max 2 years)")
    if ctx.department_id:
        department_id = ctx.department_id
    if department_id:
        d = db.get(Department, department_id)
        if d is None or d.plant_id != ctx.plant_id:
            raise HTTPException(422, "Department not found")
    return svc.Scope(plant=ctx.plant, start=start, end=end, department_id=department_id, asset_id=asset_id)


@router.get("/definitions")
def definitions(ctx: Ctx = Depends(require("kpi.view"))) -> list[dict]:
    return KPI_DEFINITIONS


@router.get("")
def all_kpis(
    ctx: Ctx = Depends(require("kpi.view")),
    db: Session = Depends(get_db, scope="function"),
    start: date | None = None,
    end: date | None = None,
    department_id: uuid.UUID | None = None,
    asset_id: uuid.UUID | None = None,
) -> dict:
    s = _scope(ctx, db, start, end, department_id, asset_id)
    return {"period": {"start": s.start, "end": s.end}, "currency": ctx.plant.currency, "kpis": svc.compute_all(db, s)}


@router.get("/trend/{code}")
def trend(
    code: str,
    ctx: Ctx = Depends(require("kpi.view")),
    db: Session = Depends(get_db, scope="function"),
    months: int = 12,
    department_id: uuid.UUID | None = None,
) -> list[dict]:
    if code not in svc.KPI_FUNCS:
        raise HTTPException(404, "Unknown KPI")
    today = plant_today(ctx.plant)
    out = []
    first = today.replace(day=1)
    manual = {  # plant-level values entered from the previous system / manual reports
        m.period_start: float(m.value)
        for m in db.execute(
            select(KPISnapshot).where(KPISnapshot.plant_id == ctx.plant_id, KPISnapshot.kpi_code == code, KPISnapshot.scope_key == MANUAL)
        ).scalars()
        if m.value is not None
    }
    for i in range(min(months, 24) - 1, -1, -1):
        y, m = divmod(first.month - 1 - i, 12)
        ms = date(first.year + y, m + 1, 1)
        me = (ms + timedelta(days=32)).replace(day=1) - timedelta(days=1)
        s = _scope(ctx, db, ms, min(me, today), department_id, None)
        r = svc.KPI_FUNCS[code](db, s)
        out.append(
            {
                "period": ms.isoformat()[:7],
                "value": None if r.value is None else round(r.value, 2),
                "available": r.available,
                "reason": r.reason,
                "manual": manual.get(ms) if not (department_id or ctx.department_id) else None,
            }
        )
    return out


MANUAL = "manual"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def _manual_values(db: Session, plant_id: uuid.UUID, year: int) -> dict[tuple[str, int], float]:
    rows = db.execute(
        select(KPISnapshot).where(
            KPISnapshot.plant_id == plant_id,
            KPISnapshot.scope_key == MANUAL,
            KPISnapshot.period_start >= date(year, 1, 1),
            KPISnapshot.period_start <= date(year, 12, 1),
        )
    ).scalars()
    return {(r.kpi_code, r.period_start.month): float(r.value) for r in rows if r.value is not None}


@router.get("/manual/values")
def manual_values(
    year: int = Query(ge=2000, le=2100), ctx: Ctx = Depends(require("kpi.view")), db: Session = Depends(get_db, scope="function")
) -> list[dict]:
    vals = _manual_values(db, ctx.plant_id, year)
    return [
        {"kpi_code": d["code"], "name": d["name"], "unit": d["unit"], "months": [vals.get((d["code"], m)) for m in range(1, 13)]}
        for d in KPI_DEFINITIONS
    ]


@router.get("/manual/template.xlsx")
def manual_template(
    year: int = Query(ge=2000, le=2100), ctx: Ctx = Depends(require("kpi.view")), db: Session = Depends(get_db, scope="function")
) -> Response:
    """One row per KPI, one column per month, pre-filled with values already entered."""
    vals = _manual_values(db, ctx.plant_id, year)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = f"KPI {year}"
    ws.append(["kpi_code", "kpi_name", "unit", *[f"{year}-{m:02d}" for m in range(1, 13)]])
    fill = PatternFill("solid", fgColor="FFF2CC")
    for d in KPI_DEFINITIONS:
        ws.append([d["code"], d["name"], d["unit"], *[vals.get((d["code"], m)) for m in range(1, 13)]])
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=15):
        for c in row:
            c.fill = fill
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 40
    ws.freeze_panes = "D2"
    h = wb.create_sheet("instructions")
    for line in (
        "Fill the yellow cells with the monthly value of each KPI (numbers only, in the unit shown in column C).",
        "Leave a cell empty to keep what is already stored; type x to delete a stored value.",
        "Do not change columns A-C or the month headers. Rows can be deleted if a KPI is not tracked.",
        "Values are stored as plant-level 'manual' figures and shown next to the values the CMMS calculates.",
    ):
        h.append([line])
    h.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="kpi_manual_{year}.xlsx"'},
    )


def _parse_manual(content: bytes) -> dict:
    try:
        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception as e:  # noqa: BLE001 - malformed spreadsheets
        raise HTTPException(422, f"Could not read file: {type(e).__name__}") from e
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it, [])]
    months: dict[int, date] = {}
    for i, h in enumerate(header):
        try:
            y, m = h.split("-")
            months[i] = date(int(y), int(m), 1)
        except ValueError:
            continue
    codes = {d["code"] for d in KPI_DEFINITIONS}
    values, clears, errors = [], [], []
    for rn, r in enumerate(it, start=2):
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        code = str(r[0]).strip() if r[0] is not None else ""
        if code not in codes:
            errors.append({"row": rn, "code": code, "errors": [f"unknown kpi_code {code or '(blank)'}"]})
            continue
        for i, ms in months.items():
            v = r[i] if i < len(r) else None
            if v is None or str(v).strip() == "":
                continue
            if str(v).strip().lower() == "x":
                clears.append((code, ms))
                continue
            try:
                values.append((code, ms, float(str(v).replace(",", "").replace("%", ""))))
            except ValueError:
                errors.append({"row": rn, "code": code, "errors": [f"{ms.isoformat()[:7]}: '{v}' is not a number"]})
    if not months:
        errors.append({"row": 1, "code": None, "errors": ["no month columns (YYYY-MM) found; use the downloaded template"]})
    return {"values": values, "clears": clears, "errors": errors}


@router.post("/manual/preview")
def manual_preview(file: UploadFile = File(...), ctx: Ctx = Depends(require("kpi.manage"))) -> dict:
    p = _parse_manual(file.file.read(5 * 1024 * 1024))
    return {"values": len(p["values"]), "cleared": len(p["clears"]), "invalid": len(p["errors"]), "errors": p["errors"]}


@router.post("/manual/commit")
def manual_commit(
    file: UploadFile = File(...), ctx: Ctx = Depends(require("kpi.manage")), db: Session = Depends(get_db, scope="function")
) -> dict:
    p = _parse_manual(file.file.read(5 * 1024 * 1024))
    if p["errors"]:
        raise HTTPException(422, {"message": "The file has errors; nothing was saved", "errors": p["errors"]})
    for code, ms in p["clears"]:
        db.execute(
            KPISnapshot.__table__.delete().where(
                KPISnapshot.plant_id == ctx.plant_id,
                KPISnapshot.kpi_code == code,
                KPISnapshot.period_start == ms,
                KPISnapshot.scope_key == MANUAL,
            )
        )
    for code, ms, v in p["values"]:
        me = (ms + timedelta(days=32)).replace(day=1) - timedelta(days=1)
        db.execute(
            pg_insert(KPISnapshot)
            .values(
                id=uuid.uuid4(),
                plant_id=ctx.plant_id,
                kpi_code=code,
                period_start=ms,
                period_end=me,
                scope_key=MANUAL,
                value=v,
                available=True,
                reason="Manual entry",
                computed_at=utcnow(),
            )
            .on_conflict_do_update(
                index_elements=["plant_id", "kpi_code", "period_start", "scope_key"],
                set_={"value": v, "available": True, "computed_at": utcnow()},
            )
        )
    audit(
        db, "kpi.manual_import", user_id=ctx.user.id, plant_id=ctx.plant_id, data={"values": len(p["values"]), "cleared": len(p["clears"])}
    )
    return {"saved": len(p["values"]), "cleared": len(p["clears"])}


@router.get("/{code}")
def one(
    code: str,
    ctx: Ctx = Depends(require("kpi.view")),
    db: Session = Depends(get_db, scope="function"),
    start: date | None = None,
    end: date | None = None,
    department_id: uuid.UUID | None = None,
    asset_id: uuid.UUID | None = None,
) -> dict:
    """Single KPI with drill-down records."""
    if code not in svc.KPI_FUNCS:
        raise HTTPException(404, "Unknown KPI")
    s = _scope(ctx, db, start, end, department_id, asset_id)
    d = next((x for x in KPI_DEFINITIONS if x["code"] == code), {})
    return {"definition": d, "period": {"start": s.start, "end": s.end}, **svc.KPI_FUNCS[code](db, s).out(with_records=True)}


class ProductionIn(In):
    department_id: uuid.UUID
    asset_id: uuid.UUID | None = None
    record_date: date
    tons_produced: float = Field(ge=0)
    good_tons: float | None = Field(None, ge=0)
    planned_production_hours: float | None = Field(None, ge=0, le=24)
    run_hours: float | None = Field(None, ge=0, le=24)
    ideal_rate_tph: float | None = Field(None, gt=0)


@router.put("/production")
def upsert_production(
    body: ProductionIn, ctx: Ctx = Depends(require("kpi.manage")), db: Session = Depends(get_db, scope="function")
) -> dict:
    d = db.get(Department, body.department_id)
    if d is None or d.plant_id != ctx.plant_id:
        raise HTTPException(422, "Department not found")
    if body.good_tons is not None and body.good_tons > body.tons_produced:
        raise HTTPException(422, "good_tons cannot exceed tons_produced")
    if body.run_hours is not None and body.planned_production_hours is not None and body.run_hours > body.planned_production_hours:
        raise HTTPException(422, "run_hours cannot exceed planned_production_hours")
    existing = db.execute(
        select(ProductionRecord).where(
            ProductionRecord.plant_id == ctx.plant_id,
            ProductionRecord.department_id == body.department_id,
            ProductionRecord.asset_id.is_(None) if body.asset_id is None else ProductionRecord.asset_id == body.asset_id,
            ProductionRecord.record_date == body.record_date,
        )
    ).scalar_one_or_none()
    if existing:
        for k, v in body.model_dump().items():
            setattr(existing, k, v)
        rec = existing
    else:
        rec = ProductionRecord(plant_id=ctx.plant_id, created_by_id=ctx.user.id, **body.model_dump())
        db.add(rec)
    db.flush()
    audit(
        db,
        "production.upsert",
        user_id=ctx.user.id,
        plant_id=ctx.plant_id,
        entity_type="production",
        entity_id=rec.id,
        data=body.model_dump(),
    )
    return {"id": rec.id}


@router.get("/production/records")
def list_production(
    ctx: Ctx = Depends(require("kpi.view")),
    db: Session = Depends(get_db, scope="function"),
    start: date | None = None,
    end: date | None = None,
) -> list[dict]:
    q = select(ProductionRecord).where(ProductionRecord.plant_id == ctx.plant_id)
    if start:
        q = q.where(ProductionRecord.record_date >= start)
    if end:
        q = q.where(ProductionRecord.record_date <= end)
    return [
        {
            "id": r.id,
            "department_id": r.department_id,
            "asset_id": r.asset_id,
            "record_date": r.record_date,
            "tons_produced": float(r.tons_produced),
            "good_tons": float(r.good_tons) if r.good_tons is not None else None,
            "planned_production_hours": r.planned_production_hours,
            "run_hours": r.run_hours,
            "ideal_rate_tph": r.ideal_rate_tph,
        }
        for r in db.execute(q.order_by(ProductionRecord.record_date.desc()).limit(2000)).scalars()
    ]


class BudgetIn(In):
    department_id: uuid.UUID | None = None
    year: int = Field(ge=2000, le=2100)
    month: int = Field(ge=1, le=12)
    amount: float = Field(ge=0)


@router.put("/budgets")
def upsert_budget(body: BudgetIn, ctx: Ctx = Depends(require("kpi.manage")), db: Session = Depends(get_db, scope="function")) -> dict:
    import uuid as _u

    if body.department_id:
        d = db.get(Department, body.department_id)
        if d is None or d.plant_id != ctx.plant_id:
            raise HTTPException(422, "Department not found")
    q = select(MaintenanceBudget).where(
        MaintenanceBudget.plant_id == ctx.plant_id,
        MaintenanceBudget.year == body.year,
        MaintenanceBudget.month == body.month,
        MaintenanceBudget.department_id.is_(None) if body.department_id is None else MaintenanceBudget.department_id == body.department_id,
    )
    b = db.execute(q).scalar_one_or_none()
    if b:
        b.amount = body.amount
    else:
        b = MaintenanceBudget(id=_u.uuid4(), plant_id=ctx.plant_id, currency=ctx.plant.currency, **body.model_dump())
        db.add(b)
    audit(db, "budget.upsert", user_id=ctx.user.id, plant_id=ctx.plant_id, data=body.model_dump())
    return {"id": b.id}


def snapshot_month(db: Session, plant, year: int, month: int) -> int:
    """Persist monthly KPI values (used by the scheduler)."""
    from app.services.common import utcnow

    ms = date(year, month, 1)
    me = (ms + timedelta(days=32)).replace(day=1) - timedelta(days=1)
    s = svc.Scope(plant=plant, start=ms, end=me)
    n = 0
    for code, f in svc.KPI_FUNCS.items():
        r = f(db, s)
        db.execute(
            pg_insert(KPISnapshot)
            .values(
                id=uuid.uuid4(),
                plant_id=plant.id,
                kpi_code=code,
                period_start=ms,
                period_end=me,
                scope_key="plant",
                value=r.value,
                available=r.available,
                reason=r.reason,
                computed_at=utcnow(),
            )
            .on_conflict_do_update(
                index_elements=["plant_id", "kpi_code", "period_start", "scope_key"],
                set_={"value": r.value, "available": r.available, "reason": r.reason, "computed_at": utcnow()},
            )
        )
        n += 1
    return n
