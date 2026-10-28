"""Roster import from the plant's own Excel grid: one row per person, one column per day, a shift code in each cell.

    Name              | 26 | 27 | 28 | ...        (or real dates in the header row)
    Mohd Wajid Khan   | D  | G  | D  |
    Alamin            | N  | N  | N  |

Several blocks may follow each other, each starting with its own "Name" header row. Codes are the plant's shift
codes (e.g. D, N, G) plus L = leave, A = absence, OFF = day off (removes an entry); a blank cell changes nothing.
Day numbers in the header are read from the chosen first month, moving to the next month when they start again
or run past the month's last day (32 = the 1st of a 31-day month's next month)."""

import io
import re
import secrets
import uuid
from datetime import date, datetime, timedelta
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.models.enums import RosterKind
from app.models.identity import Membership, Role, User
from app.models.shifts import RosterEntry, ShiftDefinition
from app.models.tenancy import Plant
from app.services.common import audit, utcnow

LEAVE = {"L", "LV", "LEAVE", "AL", "VAC", "V", "SL", "SICK"}
ABSENCE = {"A", "ABS", "ABSENT"}
OFF = {"O", "OFF", "R", "REST", "X", "-"}
MAX_DAYS = 62


def _norm(name: str) -> str:
    return re.sub(r"\s+", " ", name).strip().casefold()


def _people(db: Session, plant: Plant) -> list[User]:
    return list(
        db.execute(
            select(User)
            .join(Membership, Membership.user_id == User.id)
            .where(Membership.plant_id == plant.id, Membership.is_active.is_(True), User.is_active.is_(True))
            .order_by(User.full_name)
        ).scalars()
    )


def _shifts(db: Session, plant: Plant) -> dict[str, ShiftDefinition]:
    return {s.code.upper(): s for s in db.execute(select(ShiftDefinition).where(ShiftDefinition.plant_id == plant.id)).scalars()}


def _cell(e: RosterEntry | None, shifts: dict[uuid.UUID, str]) -> str | None:
    if e is None:
        return None
    if e.kind == RosterKind.LEAVE:
        return "L"
    if e.kind == RosterKind.ABSENCE:
        return "A"
    return shifts.get(e.shift_id) if e.shift_id else None


def template_xlsx(db: Session, plant: Plant, start: date, end: date) -> bytes:
    """The grid for a period, pre-filled with the roster already stored."""
    days = [start + timedelta(days=i) for i in range((end - start).days + 1)]
    shifts = {s.id: s.code for s in _shifts(db, plant).values()}
    stored = {
        (r.user_id, r.work_date): r
        for r in db.execute(
            select(RosterEntry).where(RosterEntry.plant_id == plant.id, RosterEntry.work_date >= start, RosterEntry.work_date <= end)
        ).scalars()
    }
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Roster"
    ws.append(["", *[d.strftime("%a").upper() for d in days]])
    ws.append(["Name", *days])
    for c in ws[2][1:]:
        c.number_format = "yyyy-mm-dd"
    for u in _people(db, plant):
        row: list[Any] = [u.full_name]
        for d in days:
            e = stored.get((u.id, d))
            row.append(_cell(e, shifts))
        ws.append(row)
    head = PatternFill("solid", fgColor="92D050")
    for r in (1, 2):
        for c in ws[r]:
            c.font = Font(bold=True)
            if c.column > 1:
                c.fill = head
    ws.column_dimensions["A"].width = 30
    ws.freeze_panes = "B3"
    h = wb.create_sheet("codes")
    h.append(["code", "meaning"])
    if not _shifts(db, plant):
        h.append(["!", "No shift types yet: add D, N, G ... in Shifts & handover › Shift types, then download this template again"])
    for s in sorted(_shifts(db, plant).values(), key=lambda s: s.start_time):
        h.append([s.code, f"{s.name} {s.start_time:%H:%M}-{s.end_time:%H:%M}, {float(s.paid_hours):g} h"])
    h.append(["L", "Leave"])
    h.append(["A", "Absence"])
    h.append(["OFF", "Day off (removes a stored entry)"])
    h.append(["(blank)", "No change"])
    h.append([])
    h.append(["Names must match the users in the CMMS (Administration › Users), or tick “create missing people” when importing."])
    h.column_dimensions["B"].width = 60
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _header_dates(row: tuple, base: date) -> dict[int, date]:
    """Column index -> date. Real dates are used as they are; bare day numbers count from the base month."""
    out: dict[int, date] = {}
    year, month, last = base.year, base.month, None
    for i, v in enumerate(row):
        if i == 0 or v is None or str(v).strip() == "":
            continue
        if isinstance(v, datetime):
            out[i] = v.date()
            continue
        if isinstance(v, date):
            out[i] = v
            continue
        txt = str(v).strip()
        m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", txt)
        if m:
            out[i] = date(int(m[1]), int(m[2]), int(m[3]))
            continue
        try:
            day = int(float(txt))
        except ValueError:
            continue
        if not 1 <= day <= MAX_DAYS:
            continue
        if last is not None and day < last:
            year, month = (year + 1, 1) if month == 12 else (year, month + 1)
        last = day
        # Past the month's last day (Excel's fill handle gives 30, 31, 32, 33 ...) runs on into the next month.
        out[i] = date(year, month, 1) + timedelta(days=day - 1)
    return out


def parse(db: Session, plant: Plant, content: bytes, base: date) -> dict[str, Any]:
    try:
        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception as e:  # noqa: BLE001 - malformed spreadsheets
        raise ValueError(f"Could not read file: {type(e).__name__}") from e
    people = {_norm(u.full_name): u for u in _people(db, plant)}
    people.update({_norm(u.email): u for u in people.copy().values()})
    shifts = _shifts(db, plant)
    cols: dict[int, date] = {}
    entries: dict[tuple[str, date], tuple[str, str | None]] = {}  # (name key, day) -> (kind, shift code)
    names: dict[str, str] = {}
    errors: list[dict[str, Any]] = []
    for rn, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue
        first = str(row[0]).strip()
        if first.casefold() in ("name", "names", "employee", "employee name"):
            cols = _header_dates(row, base)
            continue
        if not cols:
            continue
        key = _norm(first)
        names.setdefault(key, first)
        errs = []
        for i in range(1, len(row)):
            if i not in cols and row[i] is not None and str(row[i]).strip():
                errs.append(f"'{str(row[i]).strip()}' is in a column with no date in its 'Name' row; extend that header row")
                break
        for i, day in cols.items():
            v = row[i] if i < len(row) else None
            code = "" if v is None else str(v).strip().upper()
            if code == "":
                continue
            if code in shifts:
                entries[(key, day)] = (RosterKind.WORK.value, code)
            elif code in LEAVE:
                entries[(key, day)] = (RosterKind.LEAVE.value, None)
            elif code in ABSENCE:
                entries[(key, day)] = (RosterKind.ABSENCE.value, None)
            elif code in OFF:
                entries[(key, day)] = ("off", None)
            else:
                errs.append(f"{day.isoformat()}: unknown code '{code}'")
        if errs:
            errors.append({"row": rn, "code": first, "errors": errs})
    if not entries and not errors:
        errors.append({"row": 1, "code": None, "errors": ["no roster found: the header row must start with 'Name' followed by the days"]})
    days = sorted({d for _, d in entries})
    if days and (days[-1] - days[0]).days >= MAX_DAYS:
        errors.append({"row": 1, "code": None, "errors": [f"more than {MAX_DAYS} days in one file; split it by month"]})
    missing = sorted(names[k] for k in {k for k, _ in entries} if k not in people)
    return {"entries": entries, "names": names, "people": people, "shifts": shifts, "errors": errors, "missing": missing, "days": days}


def preview(p: dict[str, Any]) -> dict[str, Any]:
    persons = {k for k, _ in p["entries"]}
    return {
        "people": len(persons),
        "matched": len([k for k in persons if k in p["people"]]),
        "missing": p["missing"],
        "entries": len(p["entries"]),
        "from": p["days"][0].isoformat() if p["days"] else None,
        "to": p["days"][-1].isoformat() if p["days"] else None,
        "invalid": len(p["errors"]),
        "errors": p["errors"],
    }


def _create_person(db: Session, plant: Plant, name: str, role: Role) -> User:
    """A technician who appears on the roster but has no login yet (random password; the administrator resets it
    to give access). The e-mail is a placeholder the administrator can change."""
    slug = re.sub(r"[^a-z0-9]+", ".", name.casefold()).strip(".")[:40] or "person"
    email, n = f"{slug}@roster.local", 1
    while db.execute(select(User.id).where(User.email == email)).first():
        n += 1
        email = f"{slug}{n}@roster.local"
    u = User(
        organization_id=plant.organization_id,
        email=email,
        full_name=name,
        password_hash=hash_password(secrets.token_urlsafe(24) + "aA1!"),
        password_changed_at=utcnow(),
    )
    db.add(u)
    db.flush()
    db.add(Membership(user_id=u.id, plant_id=plant.id, role_id=role.id))
    db.flush()
    return u


def commit(db: Session, plant: Plant, p: dict[str, Any], create_missing: bool, user_id: uuid.UUID) -> dict[str, int]:
    people: dict[str, User] = dict(p["people"])
    created_people = 0
    if p["missing"]:
        if not create_missing:
            raise ValueError("Some names are not users of this plant")
        role = db.execute(select(Role).where(Role.code == "technician")).scalar_one()
        for name in p["missing"]:
            people[_norm(name)] = _create_person(db, plant, name, role)
            created_people += 1
    stored = {
        (r.user_id, r.work_date): r
        for r in db.execute(
            select(RosterEntry).where(
                RosterEntry.user_id.in_([u.id for u in people.values()]),
                RosterEntry.work_date.in_(p["days"]),
            )
        ).scalars()
    }
    saved = removed = 0
    for (key, day), (kind, code) in p["entries"].items():
        u = people[key]
        e = stored.get((u.id, day))
        if kind == "off":
            if e is not None and e.plant_id == plant.id:
                db.execute(delete(RosterEntry).where(RosterEntry.id == e.id))
                removed += 1
            continue
        if e is not None and e.plant_id != plant.id:
            continue  # rostered at another plant that day; left as it is
        s = p["shifts"].get(code) if code else None
        values = {
            "kind": RosterKind(kind),
            "shift_id": s.id if s else None,
            "hours": float(s.paid_hours) if s else 0,
            "overtime_hours": 0,
            "callout_hours": 0,
            "source": "import",
        }
        if e is None:
            db.add(RosterEntry(plant_id=plant.id, user_id=u.id, work_date=day, **values))
        else:
            for k, v in values.items():
                setattr(e, k, v)
        saved += 1
    db.flush()
    audit(
        db, "roster.import", user_id=user_id, plant_id=plant.id, data={"saved": saved, "removed": removed, "people_created": created_people}
    )
    return {"saved": saved, "removed": removed, "people_created": created_people}
