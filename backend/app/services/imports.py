"""CSV/XLSX import with preview (dry-run), validation, duplicate detection and error report.
Commit is all-or-nothing: any error aborts the whole import (safe recovery = nothing partially written)."""

import csv
import io
import uuid
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.assets import Asset
from app.models.enums import ASSET_LEVEL_ORDER, AssetLevel, Criticality
from app.models.inventory import Part
from app.models.tenancy import Department, Plant
from app.services.assets import compute_path, new_qr_token
from app.services.common import audit

MAX_ROWS = 20000

TEMPLATES: dict[str, list[tuple[str, str]]] = {
    "assets": [
        ("code", "Unique equipment code, e.g. PM1-WIRE-UTW-DR1"),
        ("name", "Description"),
        ("department_code", "PM | SP | FH | WWTP | EAI | UTIL (department code in this plant)"),
        ("parent_code", "Code of the parent asset (blank for production lines)"),
        ("level", "production_line | unit | assembly | sub_assembly | component"),
        ("criticality", "A | B | C"),
        ("oem", "Manufacturer"),
        ("model", "Model"),
        ("serial_number", "Serial"),
        ("install_date", "YYYY-MM-DD"),
        ("location", "Location"),
        ("description", "Long description"),
        ("name_ar", "Arabic name (optional)"),
    ],
    "parts": [
        ("code", "Unique item code"),
        ("description", "Description"),
        ("oem_part_number", "OEM P/N"),
        ("manufacturer", "Manufacturer"),
        ("uom", "Unit of measure, e.g. EA, KG, L"),
        ("criticality", "A | B | C"),
        ("is_critical_spare", "yes | no"),
        ("min_qty", "number"),
        ("max_qty", "number"),
        ("reorder_point", "number"),
        ("safety_stock", "number"),
        ("lead_time_days", "integer"),
        ("unit_price", "number"),
        ("category", "text"),
    ],
}


def read_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(text)))
    elif name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it, [])]
        rows = []
        for r in it:
            if r is None or all(v is None or str(v).strip() == "" for v in r):
                continue
            rows.append({header[i]: r[i] for i in range(min(len(header), len(r))) if header[i]})
    else:
        raise ValueError("Only .csv and .xlsx files are supported")
    if len(rows) > MAX_ROWS:
        raise ValueError(f"Too many rows (max {MAX_ROWS})")
    return [{(k or "").strip().lower(): (v.strip() if isinstance(v, str) else v) for k, v in r.items()} for r in rows]


def _s(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _num(v: Any, field: str, errs: list[str], integer: bool = False) -> float | int | None:
    if _s(v) is None:
        return None
    try:
        x = float(str(v).replace(",", ""))
        if x < 0:
            errs.append(f"{field} must be ≥ 0")
            return None
        return int(x) if integer else x
    except ValueError:
        errs.append(f"{field} is not a number")
        return None


def _date(v: Any, errs: list[str]) -> date | None:
    if v is None or _s(v) is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v).strip()[:10])
    except ValueError:
        errs.append("install_date must be YYYY-MM-DD")
        return None


def validate_assets(db: Session, plant: Plant, rows: list[dict]) -> dict:
    depts = {d.code.upper(): d for d in db.execute(select(Department).where(Department.plant_id == plant.id)).scalars()}
    existing = {a.code.upper(): a for a in db.execute(select(Asset).where(Asset.plant_id == plant.id)).scalars()}
    seen: dict[str, int] = {}
    file_rows: dict[str, dict] = {}
    results = []
    for i, r in enumerate(rows, start=2):  # row 1 = header
        errs: list[str] = []
        code = _s(r.get("code"))
        name = _s(r.get("name"))
        if not code:
            errs.append("code is required")
        if not name:
            errs.append("name is required")
        cu = (code or "").upper()
        if code and cu in seen:
            errs.append(f"duplicate code in file (also row {seen[cu]})")
        if code and cu in existing:
            errs.append("duplicate: code already exists in this plant")
        if code:
            seen.setdefault(cu, i)
        dept = depts.get((_s(r.get("department_code")) or "").upper())
        if dept is None:
            errs.append("unknown department_code")
        try:
            level = AssetLevel((_s(r.get("level")) or "").lower())
        except ValueError:
            errs.append("invalid level")
            level = None  # type: ignore[assignment]
        try:
            crit = Criticality((_s(r.get("criticality")) or "C").upper())
        except ValueError:
            errs.append("criticality must be A, B or C")
            crit = Criticality.C
        inst = _date(r.get("install_date"), errs)
        rec = {
            "row": i,
            "code": code,
            "name": name,
            "department": dept,
            "parent_code": _s(r.get("parent_code")),
            "level": level,
            "criticality": crit,
            "install_date": inst,
            "errors": errs,
            "raw": r,
        }
        results.append(rec)
        if code and cu not in file_rows:
            file_rows[cu] = rec
    # Hierarchy checks after all rows are known (parents may appear later in the file or already exist).
    for rec in results:
        if rec["level"] is None or rec["department"] is None:
            continue
        pc = (rec["parent_code"] or "").upper()
        if not pc:
            if rec["level"] != AssetLevel.PRODUCTION_LINE:
                rec["errors"].append("parent_code required for levels below production_line")
            continue
        parent_level, parent_dept = None, None
        if pc in existing:
            parent_level, parent_dept = existing[pc].level, existing[pc].department_id
        elif pc in file_rows and file_rows[pc] is not rec:
            parent_level = file_rows[pc]["level"]
            parent_dept = file_rows[pc]["department"].id if file_rows[pc]["department"] else None
        else:
            rec["errors"].append(f"parent_code {rec['parent_code']} not found")
            continue
        if parent_level is not None and ASSET_LEVEL_ORDER.index(rec["level"]) <= ASSET_LEVEL_ORDER.index(parent_level):
            rec["errors"].append(f"{rec['level'].value} cannot be under {parent_level.value}")
        if parent_dept is not None and parent_dept != rec["department"].id:
            rec["errors"].append("parent belongs to a different department")
    # Cycle detection among in-file parents.
    for rec in results:
        seen_chain = set()
        cur = rec
        while cur and cur["parent_code"]:
            key = cur["parent_code"].upper()
            if key in seen_chain:
                rec["errors"].append("circular parent reference")
                break
            seen_chain.add(key)
            cur = file_rows.get(key)
    errors = [{"row": r["row"], "code": r["code"], "errors": r["errors"]} for r in results if r["errors"]]
    return {"rows": results, "total": len(results), "valid": len(results) - len(errors), "invalid": len(errors), "errors": errors}


def commit_assets(db: Session, plant: Plant, rows: list[dict], user_id: uuid.UUID) -> int:
    v = validate_assets(db, plant, rows)
    if v["invalid"]:
        raise ValueError("Import has validation errors; nothing was imported")
    existing = {a.code.upper(): a for a in db.execute(select(Asset).where(Asset.plant_id == plant.id)).scalars()}
    pending = list(v["rows"])
    created: dict[str, Asset] = {}
    while pending:
        progressed = False
        for rec in list(pending):
            pc = (rec["parent_code"] or "").upper()
            parent = existing.get(pc) or created.get(pc) if pc else None
            if pc and parent is None:
                continue
            a = Asset(
                id=uuid.uuid4(),
                plant_id=plant.id,
                department_id=rec["department"].id,
                parent_id=parent.id if parent else None,
                level=rec["level"],
                code=rec["code"],
                name=rec["name"],
                name_ar=_s(rec["raw"].get("name_ar")),
                criticality=rec["criticality"],
                oem=_s(rec["raw"].get("oem")),
                model=_s(rec["raw"].get("model")),
                serial_number=_s(rec["raw"].get("serial_number")),
                install_date=rec["install_date"],
                location=_s(rec["raw"].get("location")),
                description=_s(rec["raw"].get("description")),
                qr_token=new_qr_token(),
                created_by_id=user_id,
            )
            a.path = compute_path(parent, a.id)
            db.add(a)
            created[rec["code"].upper()] = a
            pending.remove(rec)
            progressed = True
        if not progressed:
            raise ValueError("Unresolvable parent references")
    db.flush()
    audit(db, "asset.import", user_id=user_id, plant_id=plant.id, data={"created": len(created)})
    return len(created)


def validate_parts(db: Session, plant: Plant, rows: list[dict]) -> dict:
    existing = {c.upper() for c in db.execute(select(Part.code).where(Part.plant_id == plant.id)).scalars()}
    existing_pn = {
        (p.manufacturer or "").upper() + "|" + (p.oem_part_number or "").upper(): p.code
        for p in db.execute(select(Part).where(Part.plant_id == plant.id, Part.oem_part_number.is_not(None))).scalars()
    }
    seen: dict[str, int] = {}
    out = []
    for i, r in enumerate(rows, start=2):
        errs: list[str] = []
        warns: list[str] = []
        code, desc = _s(r.get("code")), _s(r.get("description"))
        if not code:
            errs.append("code is required")
        if not desc:
            errs.append("description is required")
        cu = (code or "").upper()
        if code and cu in seen:
            errs.append(f"duplicate code in file (also row {seen[cu]})")
        if code and cu in existing:
            errs.append("duplicate: code already exists")
        seen.setdefault(cu, i)
        pn_key = (_s(r.get("manufacturer")) or "").upper() + "|" + (_s(r.get("oem_part_number")) or "").upper()
        if _s(r.get("oem_part_number")) and pn_key in existing_pn:
            warns.append(f"possible duplicate of {existing_pn[pn_key]} (same manufacturer + OEM P/N)")
        try:
            crit = Criticality((_s(r.get("criticality")) or "C").upper())
        except ValueError:
            errs.append("criticality must be A, B or C")
            crit = Criticality.C
        vals = {f: _num(r.get(f), f, errs) for f in ("min_qty", "max_qty", "reorder_point", "safety_stock", "unit_price")}
        lt = _num(r.get("lead_time_days"), "lead_time_days", errs, integer=True)
        if vals["min_qty"] and vals["max_qty"] and vals["min_qty"] > vals["max_qty"]:
            errs.append("min_qty exceeds max_qty")
        crit_spare = (_s(r.get("is_critical_spare")) or "no").lower() in {"yes", "y", "true", "1"}
        out.append(
            {
                "row": i,
                "code": code,
                "description": desc,
                "criticality": crit,
                "values": vals,
                "lead_time_days": lt,
                "is_critical_spare": crit_spare,
                "raw": r,
                "errors": errs,
                "warnings": warns,
            }
        )
    errors = [{"row": r["row"], "code": r["code"], "errors": r["errors"]} for r in out if r["errors"]]
    warnings = [{"row": r["row"], "code": r["code"], "warnings": r["warnings"]} for r in out if r["warnings"]]
    return {"rows": out, "total": len(out), "valid": len(out) - len(errors), "invalid": len(errors), "errors": errors, "warnings": warnings}


def commit_parts(db: Session, plant: Plant, rows: list[dict], user_id: uuid.UUID) -> int:
    v = validate_parts(db, plant, rows)
    if v["invalid"]:
        raise ValueError("Import has validation errors; nothing was imported")
    for rec in v["rows"]:
        vals = {k: (x or 0) for k, x in rec["values"].items()}
        db.add(
            Part(
                plant_id=plant.id,
                code=rec["code"],
                description=rec["description"],
                criticality=rec["criticality"],
                oem_part_number=_s(rec["raw"].get("oem_part_number")),
                manufacturer=_s(rec["raw"].get("manufacturer")),
                uom=_s(rec["raw"].get("uom")) or "EA",
                category=_s(rec["raw"].get("category")),
                is_critical_spare=rec["is_critical_spare"],
                lead_time_days=rec["lead_time_days"],
                currency=plant.currency,
                average_cost=vals["unit_price"],
                created_by_id=user_id,
                **vals,
            )
        )
    db.flush()
    audit(db, "part.import", user_id=user_id, plant_id=plant.id, data={"created": len(v["rows"])})
    return len(v["rows"])


def template_xlsx(kind: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = kind
    cols = TEMPLATES[kind]
    ws.append([c for c, _ in cols])
    ws.append([f"<{h}>" for _, h in cols])
    help_ws = wb.create_sheet("instructions")
    help_ws.append(["column", "meaning"])
    for c, h in cols:
        help_ws.append([c, h])
    help_ws.append([])
    help_ws.append(
        [
            "Delete the example row 2 before importing. Use Preview first; nothing is written until Commit, "
            "and Commit is rejected if any row has errors."
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def errors_csv(errors: list[dict]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["row", "code", "error"])
    for e in errors:
        for msg in e["errors"]:
            w.writerow([e["row"], e["code"] or "", msg])
    return buf.getvalue()
