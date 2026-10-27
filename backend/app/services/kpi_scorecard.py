"""Company KPI scorecard: KPIs the plant reports with a monthly target and actual (e.g. a management KPI sheet).

The Excel template has two rows per KPI (Target, Actual) and one column per month, like the sheets plants already
keep. A KPI may be linked to a calculated CMMS KPI of the same department so both appear in one trend."""

import io
import uuid
from datetime import date
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.kpi_definitions import KPI_DEFINITIONS
from app.models.system import KPIScorecard, KPIScorecardValue
from app.models.tenancy import Department, Plant

COLS = ["kpi_name", "department", "unit", "higher_is_better", "cmms_kpi", "frequency", "row"]
ROWS = ("target", "actual")
YES, NO = {"yes", "y", "true", "1", "higher"}, {"no", "n", "false", "0", "lower"}


def _depts(db: Session, plant: Plant) -> dict[uuid.UUID, Department]:
    return {d.id: d for d in db.execute(select(Department).where(Department.plant_id == plant.id)).scalars()}


def scorecard(db: Session, plant: Plant, year: int) -> list[dict[str, Any]]:
    depts = _depts(db, plant)
    cards = (
        db.execute(select(KPIScorecard).where(KPIScorecard.plant_id == plant.id).order_by(KPIScorecard.sort, KPIScorecard.name))
        .scalars()
        .all()
    )
    vals: dict[tuple[uuid.UUID, int], KPIScorecardValue] = {
        (v.scorecard_id, v.period_start.month): v
        for v in db.execute(
            select(KPIScorecardValue).where(
                KPIScorecardValue.plant_id == plant.id,
                KPIScorecardValue.period_start >= date(year, 1, 1),
                KPIScorecardValue.period_start <= date(year, 12, 1),
            )
        ).scalars()
    }
    out = []
    for c in cards:
        months = []
        for m in range(1, 13):
            v = vals.get((c.id, m))
            months.append(
                {
                    "target": None if v is None or v.target is None else float(v.target),
                    "actual": None if v is None or v.actual is None else float(v.actual),
                }
            )
        d = depts.get(c.department_id) if c.department_id else None
        out.append(
            {
                "id": str(c.id),
                "name": c.name,
                "unit": c.unit,
                "department_id": str(c.department_id) if c.department_id else None,
                "department": d.code if d else None,
                "kpi_code": c.kpi_code,
                "higher_is_better": c.higher_is_better,
                "frequency": c.frequency,
                "months": months,
            }
        )
    return out


def template_xlsx(db: Session, plant: Plant, year: int) -> bytes:
    """Two rows per KPI (Target, Actual), one column per month, pre-filled with what is stored."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = f"KPI {year}"
    ws.append([*COLS, *[f"{year}-{m:02d}" for m in range(1, 13)]])
    target_fill, actual_fill = PatternFill("solid", fgColor="DDEBF7"), PatternFill("solid", fgColor="FFF2CC")
    for c in scorecard(db, plant, year):
        for kind in ROWS:
            ws.append(
                [
                    c["name"],
                    c["department"] or "",
                    c["unit"] or "",
                    "yes" if c["higher_is_better"] else "no",
                    c["kpi_code"] or "",
                    c["frequency"],
                    kind.capitalize(),
                    *[m[kind] for m in c["months"]],
                ]
            )
            for cell in ws[ws.max_row][len(COLS) :]:
                cell.fill = target_fill if kind == "target" else actual_fill
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 32
    ws.freeze_panes = "H2"
    h = wb.create_sheet("instructions")
    for line in (
        "One KPI = two rows with the same kpi_name: row 'Target' and row 'Actual'. One column per month (YYYY-MM).",
        "department: department code (blank = whole plant). unit: e.g. h, %, count, SAR/t.",
        "higher_is_better: yes (e.g. MTBF, adherence) or no (e.g. MTTR, breakdowns, cost).",
        "cmms_kpi (optional): code of the KPI the CMMS calculates, to show both in the same trend (see sheet 'codes').",
        "A blank cell keeps what is stored; type x to delete a stored value. New kpi_name rows add a KPI.",
        "Upload in KPIs & reports > Company KPIs. Nothing is saved if any row has an error.",
    ):
        h.append([line])
    h.column_dimensions["A"].width = 110
    codes = wb.create_sheet("codes")
    codes.append(["cmms_kpi", "name", "unit"])
    for d in KPI_DEFINITIONS:
        codes.append([d["code"], d["name"], d["unit"]])
    codes.append([])
    codes.append(["department", "name"])
    for d in sorted(_depts(db, plant).values(), key=lambda d: d.code):
        codes.append([d.code, d.name])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _month(h: Any) -> date | None:
    if isinstance(h, date):
        return date(h.year, h.month, 1)
    try:
        y, m = str(h).strip()[:7].split("-")
        return date(int(y), int(m), 1)
    except ValueError:
        return None


def parse(db: Session, plant: Plant, content: bytes) -> dict[str, Any]:
    try:
        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception as e:  # noqa: BLE001 - malformed spreadsheets
        raise ValueError(f"Could not read file: {type(e).__name__}") from e
    it = ws.iter_rows(values_only=True)
    header = list(next(it, []) or [])
    names = [str(h).strip().lower() if h is not None else "" for h in header]
    idx = {c: names.index(c) for c in COLS if c in names}
    months = {i: m for i, h in enumerate(header) if (m := _month(h)) is not None}
    errors: list[dict[str, Any]] = []
    if "kpi_name" not in idx or "row" not in idx or not months:
        return {
            "kpis": {},
            "errors": [{"row": 1, "code": None, "errors": ["use the downloaded template (kpi_name, row and YYYY-MM columns)"]}],
        }
    depts = {d.code.upper(): d for d in _depts(db, plant).values()}
    codes = {d["code"] for d in KPI_DEFINITIONS}
    kpis: dict[str, dict[str, Any]] = {}

    def cell(r: tuple, col: str) -> str:
        i = idx.get(col)
        v = r[i] if i is not None and i < len(r) else None
        return "" if v is None else str(v).strip()

    for rn, r in enumerate(it, start=2):
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        errs: list[str] = []
        name, kind = cell(r, "kpi_name"), cell(r, "row").lower()
        if not name:
            errs.append("kpi_name is required")
        elif len(name) > 120:
            errs.append("kpi_name is longer than 120 characters")
        if kind not in ROWS:
            errs.append("row must be Target or Actual")
        dept_code = cell(r, "department").upper()
        if dept_code and dept_code not in depts:
            errs.append(f"unknown department {dept_code}")
        link = cell(r, "cmms_kpi")
        if link and link not in codes:
            errs.append(f"unknown cmms_kpi {link}")
        hib = cell(r, "higher_is_better").lower()
        if hib and hib not in YES | NO:
            errs.append("higher_is_better must be yes or no")
        unit = cell(r, "unit")
        if len(unit) > 20:
            errs.append("unit is longer than 20 characters")
        k = kpis.setdefault(
            name.lower(),
            {"name": name, "rows": set(), "values": [], "clears": [], "order": len(kpis), "first_row": rn},
        )
        if kind in k["rows"]:
            errs.append(f"{name}: a second '{kind}' row")
        k["rows"].add(kind)
        meta = {
            "department": depts[dept_code] if dept_code in depts else None,
            "unit": unit or None,
            "kpi_code": link or None,
            "higher_is_better": hib not in NO,
            "frequency": cell(r, "frequency") or "Monthly",
        }
        if "meta" in k and k["meta"]["department"] != meta["department"]:
            errs.append(f"{name}: Target and Actual rows have different departments")
        k.setdefault("meta", meta)
        for i, ms in months.items():
            v = r[i] if i < len(r) else None
            if v is None or str(v).strip() == "" or kind not in ROWS:
                continue
            if str(v).strip().lower() == "x":
                k["clears"].append((kind, ms))
                continue
            try:
                k["values"].append((kind, ms, float(str(v).replace(",", "").replace("%", ""))))
            except ValueError:
                errs.append(f"{ms.isoformat()[:7]}: '{v}' is not a number")
        if errs:
            errors.append({"row": rn, "code": name or None, "errors": errs})
    return {"kpis": kpis, "errors": errors}


def preview(p: dict[str, Any]) -> dict[str, Any]:
    ks = p["kpis"].values()
    return {
        "kpis": len(ks),
        "values": sum(len(k["values"]) for k in ks),
        "cleared": sum(len(k["clears"]) for k in ks),
        "invalid": len(p["errors"]),
        "errors": p["errors"],
    }


def commit(db: Session, plant: Plant, p: dict[str, Any]) -> dict[str, int]:
    existing = {c.name.lower(): c for c in db.execute(select(KPIScorecard).where(KPIScorecard.plant_id == plant.id)).scalars()}
    created = saved = cleared = 0
    for key, k in p["kpis"].items():
        m = k["meta"]
        c = existing.get(key)
        if c is None:
            c = KPIScorecard(plant_id=plant.id, name=k["name"])
            db.add(c)
            created += 1
        c.department_id = m["department"].id if m["department"] else None
        c.unit, c.kpi_code, c.higher_is_better, c.frequency, c.sort = (
            m["unit"],
            m["kpi_code"],
            m["higher_is_better"],
            m["frequency"],
            k["order"],
        )
        db.flush()
        rows = {v.period_start: v for v in db.execute(select(KPIScorecardValue).where(KPIScorecardValue.scorecard_id == c.id)).scalars()}
        for kind, ms, val in k["values"]:
            v = rows.get(ms)
            if v is None:
                v = rows[ms] = KPIScorecardValue(plant_id=plant.id, scorecard_id=c.id, period_start=ms)
                db.add(v)
            setattr(v, kind, val)
            saved += 1
        for kind, ms in k["clears"]:
            if (v := rows.get(ms)) is not None:
                setattr(v, kind, None)
                cleared += 1
        db.flush()
        for v in rows.values():
            if v.target is None and v.actual is None and v.id is not None:
                db.delete(v)
    return {"kpis_created": created, "saved": saved, "cleared": cleared}
