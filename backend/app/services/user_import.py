"""Users of a plant from Excel: add many people at once, or change role, department, trade or e-mail of existing ones.

The template comes filled with the plant's current users. A row whose old_email matches a user updates that user
(email may be changed, e.g. a roster placeholder name@roster.local to the real address); a row without old_email
adds a person. New people, and anyone marked reset_password, get a generated temporary password that is returned
once, to the administrator who imports, to hand out. Preview first; commit is all-or-nothing."""

import io
import re
import secrets
import uuid
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select, update
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.models.identity import Membership, Role, User, UserSession
from app.models.tenancy import Department, Plant
from app.services.common import audit, utcnow
from app.services.imports import _s, read_rows

COLS = ["full_name", "email", "role", "department_code", "trade", "employee_no", "hourly_rate", "reset_password", "old_email"]
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_ALPHABET = "abcdefghjkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789"


def temp_password() -> str:
    """Three groups of four, e.g. Kx7m-Qp2r-Zt9w: easy to read out, and meets the password policy."""
    while True:
        groups = ["".join(secrets.choice(_ALPHABET) for _ in range(4)) for _ in range(3)]
        pw = "-".join(groups)
        if any(c.islower() for c in pw) and any(c.isupper() for c in pw) and any(c.isdigit() for c in pw):
            return pw


def _roles(db: Session) -> dict[str, Role]:
    out: dict[str, Role] = {}
    for r in db.execute(select(Role)).scalars():
        out[r.code.lower()] = r
        out[r.name.lower()] = r
    return out


def _members(db: Session, plant: Plant) -> list[tuple[Membership, User]]:
    return list(
        db.execute(select(Membership, User).join(User, User.id == Membership.user_id).where(Membership.plant_id == plant.id)).tuples()
    )


def template_xlsx(db: Session, plant: Plant) -> bytes:
    depts = {d.id: d.code for d in db.execute(select(Department).where(Department.plant_id == plant.id)).scalars()}
    roles = sorted({(r.code, r.name) for r in _roles(db).values()})
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "users"
    ws.append(COLS)
    for m, u in sorted(_members(db, plant), key=lambda x: x[1].full_name.lower()):
        if not (m.is_active and u.is_active):
            continue
        ws.append(
            [
                u.full_name,
                u.email,
                m.role.code,
                depts.get(m.department_id, ""),
                u.trade,
                u.employee_no,
                float(u.hourly_rate) if u.hourly_rate is not None else None,
                "",
                u.email,
            ]
        )
    head = PatternFill("solid", fgColor="92D050")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head
    for col, w in zip("ABCDEFGHI", (28, 32, 22, 16, 14, 13, 12, 14, 32), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    last = max(ws.max_row + 300, 400)
    dv_role = DataValidation(type="list", formula1=f'"{",".join(c for c, _ in roles)}"', allow_blank=False)
    dv_dept = DataValidation(type="list", formula1=f'"{",".join(sorted(depts.values()))}"', allow_blank=True)
    dv_yes = DataValidation(type="list", formula1='"yes"', allow_blank=True)
    for dv, col in ((dv_role, "C"), (dv_dept, "D"), (dv_yes, "H")):
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
    h = wb.create_sheet("help")
    for row in (
        ["full_name", "Name as it should appear in the system"],
        ["email", "Login e-mail (unique). To change someone's e-mail, type the new one here and keep the old one in old_email"],
        ["role", "One of: " + ", ".join(f"{c} ({n})" for c, n in roles)],
        ["department_code", "Limits the person to that area (" + ", ".join(sorted(depts.values())) + "); blank = whole plant"],
        ["trade", "e.g. mechanical, electrical, instrument (optional)"],
        ["employee_no / hourly_rate", "Optional; hourly rate is used for labour cost"],
        ["reset_password", "yes = give an existing user a new temporary password"],
        ["old_email", "Filled in for existing users: do not change it. Leave blank on new rows."],
        [],
        ["New people get a temporary password, shown once after the import: hand it to them; they change it after signing in."],
        ["Rows of people you do not want to change can stay as they are, or be deleted from the file: nobody is removed by an import."],
    ):
        h.append(row)
    h.column_dimensions["A"].width = 24
    h.column_dimensions["B"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def check(db: Session, plant: Plant, filename: str, content: bytes, may_grant_admin: bool) -> dict[str, Any]:
    rows = read_rows(filename, content)
    roles = _roles(db)
    depts = {d.code.upper(): d for d in db.execute(select(Department).where(Department.plant_id == plant.id)).scalars()}
    members = {u.email.lower(): (m, u) for m, u in _members(db, plant)}
    all_users = {u.email.lower(): u for u in db.execute(select(User).where(User.organization_id == plant.organization_id)).scalars()}
    plan: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: dict[str, int] = {}
    unchanged = 0
    for i, r in enumerate(rows, start=2):
        name, email, old = _s(r.get("full_name")), (_s(r.get("email")) or "").lower(), (_s(r.get("old_email")) or "").lower()
        if not name and not email:
            continue
        errs: list[str] = []
        if not name or len(name) < 2:
            errs.append("full_name is required")
        if not EMAIL.match(email):
            errs.append("email is not a valid e-mail address")
        role = roles.get((_s(r.get("role")) or "").lower())
        if role is None:
            errs.append("unknown role")
        elif role.code == "system_admin" and not may_grant_admin:
            errs.append("only a system administrator can give the system administrator role")
        dcode = (_s(r.get("department_code")) or "").upper()
        dept = depts.get(dcode) if dcode else None
        if dcode and dept is None:
            errs.append(f"department {dcode} not found")
        rate = None
        if _s(r.get("hourly_rate")) is not None:
            try:
                rate = float(str(r.get("hourly_rate")).replace(",", ""))
                if rate < 0:
                    raise ValueError
            except ValueError:
                errs.append("hourly_rate must be a number ≥ 0")
        reset = (_s(r.get("reset_password")) or "").lower() in ("yes", "y", "1", "true", "نعم")
        current = members.get(old) if old else members.get(email)
        if old and current is None:
            errs.append(f"old_email {old} is not a user of this plant")
        if email in seen:
            errs.append(f"e-mail also on row {seen[email]}")
        seen.setdefault(email, i)
        other = all_users.get(email)
        if current is not None and other is not None and other.id != current[1].id:
            errs.append(f"{email} is already used by another user")
        if errs:
            errors.append({"row": i, "code": email or name, "errors": errs})
            continue
        assert role is not None and name is not None
        vals = {
            "full_name": name,
            "email": email,
            "role": role,
            "dept": dept,
            "trade": _s(r.get("trade")),
            "employee_no": _s(r.get("employee_no")),
            "hourly_rate": rate,
            "reset": reset,
        }
        if current is None:
            plan.append({"action": "add", **vals, "existing_user": other})
            continue
        m, u = current
        same = (
            u.full_name == name
            and u.email.lower() == email
            and m.role_id == role.id
            and m.department_id == (dept.id if dept else None)
            and (u.trade or None) == vals["trade"]
            and (u.employee_no or None) == vals["employee_no"]
            and (float(u.hourly_rate) if u.hourly_rate is not None else None) == rate
            and m.is_active
            and u.is_active
            and not reset
        )
        if same:
            unchanged += 1
        else:
            plan.append({"action": "update", **vals, "membership": m, "user": u})
    adds = [p for p in plan if p["action"] == "add"]
    return {
        "plan": plan,
        "preview": {
            "new": len(adds),
            "updated": len(plan) - len(adds),
            "unchanged": unchanged,
            "passwords": len(adds) + sum(1 for p in plan if p["action"] == "update" and p["reset"]),
            "invalid": len(errors),
            "errors": errors,
        },
    }


def commit(db: Session, plant: Plant, checked: dict[str, Any], admin_id: uuid.UUID) -> dict[str, Any]:
    if checked["preview"]["invalid"]:
        raise ValueError("The file has errors; nothing was imported")
    now = utcnow()
    credentials: list[dict[str, str]] = []
    for p in checked["plan"]:
        dept_id = p["dept"].id if p["dept"] else None
        if p["action"] == "add":
            pw = temp_password()
            u = p["existing_user"]
            if u is None:
                u = User(
                    organization_id=plant.organization_id,
                    email=p["email"],
                    full_name=p["full_name"],
                    password_hash=hash_password(pw),
                    password_changed_at=now,
                )
                db.add(u)
            else:  # someone removed earlier, or a member of another plant of the organization
                u.full_name, u.is_active = p["full_name"], True
                u.password_hash, u.password_changed_at = hash_password(pw), now
            u.trade, u.employee_no, u.hourly_rate = p["trade"], p["employee_no"], p["hourly_rate"]
            db.flush()
            m = db.execute(select(Membership).where(Membership.user_id == u.id, Membership.plant_id == plant.id)).scalar_one_or_none()
            if m is None:
                db.add(Membership(user_id=u.id, plant_id=plant.id, role_id=p["role"].id, department_id=dept_id))
            else:
                m.role_id, m.department_id, m.is_active = p["role"].id, dept_id, True
            credentials.append({"name": p["full_name"], "email": p["email"], "temporary_password": pw})
            continue
        u, m = p["user"], p["membership"]
        u.full_name, u.email = p["full_name"], p["email"]
        u.trade, u.employee_no, u.hourly_rate = p["trade"], p["employee_no"], p["hourly_rate"]
        u.is_active, m.is_active = True, True
        m.role_id, m.department_id = p["role"].id, dept_id
        if p["reset"]:
            pw = temp_password()
            u.password_hash, u.password_changed_at = hash_password(pw), now
            db.execute(update(UserSession).where(UserSession.user_id == u.id, UserSession.revoked_at.is_(None)).values(revoked_at=now))
            credentials.append({"name": p["full_name"], "email": p["email"], "temporary_password": pw})
    db.flush()
    pv = checked["preview"]
    audit(
        db,
        "user.import",
        user_id=admin_id,
        plant_id=plant.id,
        data={"new": pv["new"], "updated": pv["updated"], "passwords": len(credentials)},
    )
    return {"created": pv["new"], "changed": pv["updated"], "credentials": credentials}
