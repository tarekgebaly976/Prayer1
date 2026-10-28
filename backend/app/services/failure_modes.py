"""Failure-mode catalogue organised by equipment type (motors, pumps, gearboxes ...), as used for RCM/FMEA.

- Failure codes carry an equipment type and a category (Mechanical / Electrical / Instrument), imported from Excel.
- Assets carry an equipment type, set by hand, by the asset import, or suggested from the asset name.
- A work order offers the failure modes of its equipment type first (the asset's own type, else the nearest
  parent's), so the technician picks from a short list.
- The Pareto report counts work orders per failure mode, with downtime, for Top-10 / RCA reviews."""

import io
import re
import uuid
from datetime import date, datetime, time, timedelta
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.assets import Asset, FailureCode
from app.models.enums import FailureCodeKind, WOStatus, WOType
from app.models.tenancy import Plant
from app.models.workorders import WorkOrder
from app.services.common import audit
from app.services.imports import _s, read_rows

# Equipment types with the words that identify them in an asset name. Checked in this order: the first match wins,
# so "Pump motor" is a motor and "MCC panel" is switchgear.
TYPE_WORDS: list[tuple[str, tuple[str, ...]]] = [
    ("Transformers", ("transformer",)),
    ("Switchgear / MCC", ("mcc", "switchgear", "circuit breaker", "distribution board", "switchboard")),
    ("PLC / HMI / Automation", ("plc", "hmi", "dcs", "scada", "i/o module")),
    (
        "Instrumentation",
        (
            "transmitter",
            "sensor",
            "gauge",
            "flowmeter",
            "flow meter",
            "thermocouple",
            "rtd",
            "analyzer",
            "analyser",
            "level switch",
            "pressure switch",
            "probe",
        ),
    ),
    ("Control Valves", ("control valve", "valve")),
    ("Electric Motors", ("motor",)),
    ("Gearboxes", ("gearbox", "gear box", "reducer", "gear unit")),
    ("Compressors", ("compressor",)),
    ("Pumps", ("pump",)),
    ("Fans & Blowers", ("fan", "blower")),
    ("Conveyors", ("conveyor",)),
    ("Piping Systems", ("pipe", "piping", "pipeline")),
    # Paper-machine equipment (second catalogue). After the general types, so "Felt roll drive motor" stays a motor.
    ("Cranes & Hoists", ("crane", "autocrane", "hoist", "winch")),
    ("Hydraulic & Pneumatic Units", ("hydraulic", "pneumatic", "power pack", "hpu", "air cylinder")),
    ("Doctors & Showers", ("doctor", "shower", "showers")),
    ("Headboxes", ("headbox", "headboxes", "head box")),
    ("Agitators & Mixers", ("agitator", "mixer")),
    ("Refiners & Pulpers", ("refiner", "pulper", "deflaker", "disperger", "disperser")),
    ("Screens & Cleaners", ("screen", "cleaner", "centricleaner", "fractionator")),
    ("Heat Exchangers", ("heat exchanger", "condenser", "steam trap")),
    ("Rolls", ("roll", "roller", "rolls", "rollers")),
    ("Dryer Cylinders", ("cylinder", "yankee", "dryer can", "drying cylinder")),
    ("Felts & Fabrics", ("felt", "fabric", "forming fabric", "press felt")),
    ("Tanks & Chests", ("tank", "chest", "silo", "vessel")),
]
EQUIPMENT_TYPES = [t for t, _ in TYPE_WORDS]
_PATTERNS = [(t, re.compile(r"\b(" + "|".join(re.escape(w) for w in words) + r")s?\b", re.I)) for t, words in TYPE_WORDS]
COLS = ["kind", "code", "description", "equipment_type", "category", "description_ar"]


def guess_type(name: str) -> str | None:
    for t, pat in _PATTERNS:
        if pat.search(name):
            return t
    return None


def classify(db: Session, plant: Plant, apply: bool, user_id: uuid.UUID | None = None) -> dict[str, Any]:
    """Suggest an equipment type from the name for assets that have none; apply=True saves the suggestions."""
    assets = db.execute(select(Asset).where(Asset.plant_id == plant.id, Asset.equipment_type.is_(None))).scalars().all()
    by_type: dict[str, int] = {}
    samples: dict[str, list[str]] = {}
    unmatched = 0
    for a in assets:
        t = guess_type(a.name)
        if t is None:
            unmatched += 1
            continue
        by_type[t] = by_type.get(t, 0) + 1
        if len(samples.setdefault(t, [])) < 5:
            samples[t].append(f"{a.code} — {a.name}")
        if apply:
            a.equipment_type = t
    if apply:
        db.flush()
        audit(db, "asset.classify", user_id=user_id, plant_id=plant.id, data={"set": sum(by_type.values())})
    return {"set": sum(by_type.values()), "by_type": by_type, "samples": samples, "unmatched": unmatched}


def asset_type(db: Session, asset: Asset) -> str | None:
    """The asset's own equipment type, else the nearest parent's (a bearing under a pump gets the pump's modes)."""
    if asset.equipment_type:
        return asset.equipment_type
    ids = [uuid.UUID(x) for x in asset.path.strip("/").split("/") if x][:-1]
    if not ids:
        return None
    types = dict(db.execute(select(Asset.id, Asset.equipment_type).where(Asset.id.in_(ids))).tuples().all())
    for i in reversed(ids):
        if types.get(i):
            return types[i]
    return None


def export_xlsx(db: Session, plant: Plant) -> bytes:
    """The current catalogue in the import layout (also the template: add rows, upload)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Failure codes"
    ws.append(COLS)
    for c in db.execute(
        select(FailureCode).where(FailureCode.plant_id == plant.id).order_by(FailureCode.kind, FailureCode.equipment_type, FailureCode.code)
    ).scalars():
        ws.append([c.kind.value, c.code, c.description, c.equipment_type, c.category, c.description_ar])
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEF", (10, 16, 40, 26, 14, 30), strict=True):
        ws.column_dimensions[col].width = w
    h = wb.create_sheet("help")
    for row in (
        ["kind", "problem (failure mode) | cause | remedy"],
        ["code", "Unique per kind, e.g. PMP-M02 (letters, digits, - _ . /)"],
        ["description", "e.g. Mechanical seal failure"],
        ["equipment_type", "One of: " + ", ".join(EQUIPMENT_TYPES) + " (or your own; blank = any equipment)"],
        ["category", "Mechanical | Electrical | Instrument (optional)"],
        ["description_ar", "Arabic text (optional)"],
        [],
        ["An existing code is updated with the new text; codes not in the file are left as they are."],
    ):
        h.append(row)
    h.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cut(v: Any, n: int) -> str | None:
    t = _s(v)
    return t[:n] if t else None


def check(db: Session, plant: Plant, filename: str, content: bytes) -> dict[str, Any]:
    rows = read_rows(filename, content)
    existing = {(c.kind, c.code.upper()): c for c in db.execute(select(FailureCode).where(FailureCode.plant_id == plant.id)).scalars()}
    seen: dict[tuple, int] = {}
    good: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for i, r in enumerate(rows, start=2):
        errs: list[str] = []
        try:
            kind = FailureCodeKind((_s(r.get("kind")) or "problem").lower())
        except ValueError:
            errs.append("kind must be problem, cause or remedy")
            kind = FailureCodeKind.PROBLEM
        code = _s(r.get("code"))
        desc = _s(r.get("description"))
        if not code or not re.fullmatch(r"[A-Za-z0-9._\-/]{1,32}", code):
            errs.append("code is required: up to 32 letters, digits or - _ . /")
        if not desc:
            errs.append("description is required")
        key = (kind, (code or "").upper())
        if code and key in seen:
            errs.append(f"duplicate code in file (also row {seen[key]})")
        seen.setdefault(key, i)
        if errs:
            errors.append({"row": i, "code": code, "errors": errs})
            continue
        cat = _cut(r.get("category"), 30)
        good.append(
            {
                "key": key,
                "kind": kind,
                "code": code,
                "description": (desc or "")[:200],
                "equipment_type": _cut(r.get("equipment_type"), 60),
                "category": cat.title() if cat else None,
                "description_ar": _cut(r.get("description_ar"), 200),
            }
        )
    new = sum(1 for g in good if g["key"] not in existing)
    return {
        "rows": good,
        "preview": {
            "new": new,
            "updated": len(good) - new,
            "invalid": len(errors),
            "errors": errors,
            "types": sorted({g["equipment_type"] for g in good if g["equipment_type"]}),
        },
    }


def commit(db: Session, plant: Plant, checked: dict[str, Any], user_id: uuid.UUID) -> dict[str, int]:
    if checked["preview"]["invalid"]:
        raise ValueError("The file has errors; nothing was imported")
    existing = {(c.kind, c.code.upper()): c for c in db.execute(select(FailureCode).where(FailureCode.plant_id == plant.id)).scalars()}
    for g in checked["rows"]:
        vals = {k: g[k] for k in ("description", "equipment_type", "category", "description_ar")}
        c = existing.get(g["key"])
        if c is None:
            db.add(FailureCode(plant_id=plant.id, kind=g["kind"], code=g["code"], **vals))
        else:
            for k, v in vals.items():
                setattr(c, k, v)
            c.is_active = True
    db.flush()
    p = checked["preview"]
    audit(db, "failure_code.import", user_id=user_id, plant_id=plant.id, data={"new": p["new"], "updated": p["updated"]})
    return {"new": p["new"], "updated": p["updated"]}


def pareto(db: Session, plant: Plant, start: date, end: date, under: Asset | None, by: str) -> dict[str, Any]:
    """Work orders per failure mode (problem code) in the period, most frequent first. by='mode' adds up the same
    failure mode across equipment types (e.g. all bearing failures); by='code' keeps them apart; by='symptom'
    counts the reported symptoms instead."""
    code_col = WorkOrder.symptom_code_id if by == "symptom" else WorkOrder.problem_code_id
    tz = ZoneInfo(plant.timezone)
    t0 = datetime.combine(start, time.min, tz)
    t1 = datetime.combine(end + timedelta(days=1), time.min, tz)
    when = func.coalesce(WorkOrder.actual_start, WorkOrder.created_at)
    q = (
        select(
            FailureCode.code,
            FailureCode.description,
            FailureCode.equipment_type,
            func.count(WorkOrder.id),
            func.coalesce(func.sum(WorkOrder.downtime_hours), 0),
            func.max(when),
        )
        .join(FailureCode, FailureCode.id == code_col)
        .where(WorkOrder.plant_id == plant.id, WorkOrder.status != WOStatus.CANCELLED, when >= t0, when < t1)
        .group_by(FailureCode.code, FailureCode.description, FailureCode.equipment_type)
    )
    if under is not None:
        q = q.join(Asset, Asset.id == WorkOrder.asset_id).where(Asset.path.startswith(under.path))
    rows: dict[str, dict[str, Any]] = {}
    for code, desc, etype, n, dt, last in db.execute(q).tuples():
        key = desc.strip().casefold() if by == "mode" else code
        r = rows.setdefault(
            key,
            {
                "code": None if by == "mode" else code,
                "failure_mode": desc,
                "equipment_types": set(),
                "work_orders": 0,
                "downtime_hours": 0.0,
                "last": None,
            },
        )
        if etype:
            r["equipment_types"].add(etype)
        r["work_orders"] += n
        r["downtime_hours"] += float(dt)
        r["last"] = max(filter(None, [r["last"], last])) if last else r["last"]
    total = sum(r["work_orders"] for r in rows.values())
    out = sorted(rows.values(), key=lambda r: (-r["work_orders"], -r["downtime_hours"], r["failure_mode"]))
    cum = 0
    for r in out:
        cum += r["work_orders"]
        r["equipment_types"] = sorted(r["equipment_types"])
        r["share_pct"] = round(100 * r["work_orders"] / total, 1) if total else 0
        r["cumulative_pct"] = round(100 * cum / total, 1) if total else 0
        r["last"] = r["last"].isoformat() if r["last"] else None
    without = db.execute(
        select(func.count(WorkOrder.id)).where(
            WorkOrder.plant_id == plant.id,
            WorkOrder.problem_code_id.is_(None),
            WorkOrder.status == WOStatus.CLOSED,
            when >= t0,
            when < t1,
            WorkOrder.type.in_((WOType.CORRECTIVE, WOType.EMERGENCY)),
        )
    ).scalar_one()
    return {"rows": out, "total": total, "closed_breakdowns_without_code": without}
