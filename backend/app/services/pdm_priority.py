"""Measurement-point priorities in bulk: download the point list as Excel, set the priority column, upload it.

Priority sets the list order technicians see (1 critical, 2 high, 3 normal). Only the priority column is read;
every other column is there to recognise the point. Preview first; apply is all-or-nothing."""

import io
import uuid
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.assets import Asset
from app.models.pdm import MeasurementPoint
from app.models.tenancy import Plant
from app.services.common import audit
from app.services.imports import _s, read_rows

NAMES = {1: "Critical", 2: "High", 3: "Normal"}
_WORDS = {"1": 1, "critical": 1, "c": 1, "2": 2, "high": 2, "h": 2, "3": 3, "normal": 3, "n": 3, "": 3}


def parse_priority(v: Any) -> int | None:
    """1/2/3 or critical/high/normal (any case); blank = normal. None if not recognised."""
    txt = (_s(v) or "").lower()
    if txt.endswith(".0"):
        txt = txt[:-2]
    return _WORDS.get(txt)


def export_xlsx(db: Session, plant: Plant, points: list[MeasurementPoint]) -> bytes:
    codes = dict(db.execute(select(Asset.id, Asset.code).where(Asset.id.in_({p.asset_id for p in points}))).tuples().all())
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Priorities"
    ws.append(["point_code", "priority", "asset_code", "point_name", "technique", "last_status"])
    for p in points:
        ws.append([p.code, NAMES.get(p.priority, "Normal"), codes.get(p.asset_id, ""), p.name, p.technique.value, p.last_status.value])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws["B1"].fill = PatternFill("solid", fgColor="FFC000")  # the only column that is read
    dv = DataValidation(type="list", formula1='"Critical,High,Normal"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{max(ws.max_row, 2)}")
    for col, w in zip("ABCDEF", (24, 11, 22, 60, 14, 12), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def check(db: Session, plant: Plant, filename: str, content: bytes, scope: Any = True) -> dict[str, Any]:
    """`scope`: SQL condition limiting which points may be changed (the user's department)."""
    rows = read_rows(filename, content)
    points = {p.code.upper(): p for p in db.execute(select(MeasurementPoint).where(MeasurementPoint.plant_id == plant.id, scope)).scalars()}
    changes: dict[uuid.UUID, int] = {}
    errors: list[dict[str, Any]] = []
    unchanged = 0
    for i, r in enumerate(rows, start=2):
        code = _s(r.get("point_code"))
        if not code:
            continue
        p = points.get(code.upper())
        prio = parse_priority(r.get("priority"))
        errs = []
        if p is None:
            errs.append(f"measurement point {code} not found")
        if prio is None:
            errs.append("priority must be Critical, High or Normal (or 1, 2, 3)")
        if errs:
            errors.append({"row": i, "code": code, "errors": errs})
        elif p is not None and prio is not None:
            if p.priority == prio:
                unchanged += 1
            else:
                changes[p.id] = prio
    if not rows or ("point_code" not in rows[0] and not errors):
        errors.append({"row": 1, "code": None, "errors": ["the first row must have the columns point_code and priority"]})
    by_prio = {NAMES[n]: sum(1 for v in changes.values() if v == n) for n in NAMES}
    return {
        "changes": changes,
        "preview": {"changed": len(changes), "unchanged": unchanged, "to": by_prio, "invalid": len(errors), "errors": errors},
    }


def apply(db: Session, plant: Plant, checked: dict[str, Any], user_id: uuid.UUID) -> dict[str, int]:
    if checked["preview"]["invalid"]:
        raise ValueError("The file has errors; nothing was changed")
    for pid, prio in checked["changes"].items():
        p = db.get(MeasurementPoint, pid)
        assert p is not None and p.plant_id == plant.id
        p.priority = prio
    db.flush()
    audit(db, "pdm.priority_import", user_id=user_id, plant_id=plant.id, data={"changed": len(checked["changes"])})
    return {"changed": len(checked["changes"])}
