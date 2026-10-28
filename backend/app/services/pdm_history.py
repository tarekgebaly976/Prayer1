"""Condition-monitoring history from Excel: one row per reading (point_code, date, optional time, value, note).

Readings get their status from the point's current limits, but no alerts or notifications are raised for old
values. A point's "last value" moves only when an imported reading is newer than the one it has. Re-importing the
same file adds nothing (a point + time is stored once per source). Preview first; commit is all-or-nothing."""

import io
import uuid
from datetime import date, datetime, time, timedelta
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select, tuple_
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

from app.models.assets import Asset
from app.models.enums import ConditionStatus
from app.models.pdm import ConditionReading, MeasurementPoint
from app.models.tenancy import Plant
from app.services.common import audit, utcnow
from app.services.imports import _date, _s, read_rows
from app.services.pdm import evaluate

COLS = ["point_code", "date", "time", "value", "note"]
SOURCE = "import"
DEFAULT_TIME = time(8, 0)


def template_xlsx(db: Session, plant: Plant, points: list[MeasurementPoint]) -> bytes:
    """The import layout plus a sheet listing every point code, to look codes up while filling it."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Readings"
    ws.append(COLS)
    if points:
        ws.append([points[0].code, date.today().isoformat(), "08:00", None, "example row: delete it"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="92D050")
    for col, w in zip("ABCDE", (26, 12, 8, 10, 40), strict=True):
        ws.column_dimensions[col].width = w
    codes = dict(db.execute(select(Asset.id, Asset.code).where(Asset.id.in_({p.asset_id for p in points}))).tuples().all())
    pl = wb.create_sheet("points")
    pl.append(["point_code", "asset_code", "point_name", "technique", "unit", "warning_high", "alarm_high"])
    for p in points:
        pl.append(
            [
                p.code,
                codes.get(p.asset_id, ""),
                p.name,
                p.technique.value,
                p.unit,
                float(p.warning_high) if p.warning_high is not None else None,
                float(p.alarm_high) if p.alarm_high is not None else None,
            ]
        )
    for c in pl[1]:
        c.font = Font(bold=True)
    pl.column_dimensions["A"].width = 26
    pl.column_dimensions["C"].width = 70
    h = wb.create_sheet("help")
    for row in (
        ["point_code", "Code of an existing measurement point (see the 'points' sheet)"],
        ["date", "YYYY-MM-DD (or an Excel date)"],
        ["time", "Optional HH:MM; blank = 08:00. Two readings of a point on one day need different times."],
        ["value", "The reading, in the point's unit"],
        ["note", "Optional, e.g. 'Day shift' or who measured"],
        [],
        ["Old readings get a status from today's limits but raise no alerts. Importing the same file twice adds nothing."],
    ):
        h.append(row)
    h.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _time(v: Any, errs: list[str]) -> time:
    if v is None or str(v).strip() == "":
        return DEFAULT_TIME
    if isinstance(v, time):
        return v
    if isinstance(v, datetime):
        return v.time()
    try:
        hh, mm = str(v).strip().split(":")[:2]
        return time(int(hh), int(mm))
    except ValueError:
        errs.append("time must be HH:MM")
        return DEFAULT_TIME


def check(db: Session, plant: Plant, filename: str, content: bytes, scope: Any = True) -> dict[str, Any]:
    """`scope`: SQL condition limiting which points may be written (the user's department)."""
    rows = read_rows(filename, content)
    points = {p.code.upper(): p for p in db.execute(select(MeasurementPoint).where(MeasurementPoint.plant_id == plant.id, scope)).scalars()}
    tz = ZoneInfo(plant.timezone)
    now = utcnow()
    good: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: dict[tuple, int] = {}
    for i, r in enumerate(rows, start=2):
        code = _s(r.get("point_code"))
        if not code and _s(r.get("value")) is None:
            continue  # empty line
        errs: list[str] = []
        p = points.get((code or "").upper())
        if not code:
            errs.append("point_code is blank")
        elif p is None:
            errs.append(f"measurement point {code} not found")
        d = _date(r.get("date"), errs, "date")
        if d is None and not errs:
            errs.append("date is required")
        at = datetime.combine(d, _time(r.get("time"), errs), tz) if d else None
        if at is not None and at > now + timedelta(minutes=10):
            errs.append("date is in the future")
        try:
            value = float(str(r.get("value")).replace(",", "."))
        except (TypeError, ValueError):
            errs.append("value is not a number")
            value = 0.0
        if p is not None and at is not None:
            key = (p.id, at)
            if key in seen:
                errs.append(f"same point and time as row {seen[key]}")
            seen.setdefault(key, i)
        if errs:
            errors.append({"row": i, "code": code, "errors": errs})
            continue
        assert p is not None and at is not None
        good.append({"point": p, "at": at, "value": value, "note": _s(r.get("note")), "status": evaluate(p, value)[0]})
    existing = set()
    if good:
        keys = [(g["point"].id, g["at"]) for g in good]
        existing = set(
            db.execute(
                select(ConditionReading.point_id, ConditionReading.measured_at).where(
                    ConditionReading.source == SOURCE, tuple_(ConditionReading.point_id, ConditionReading.measured_at).in_(keys)
                )
            ).tuples()
        )
    new = [g for g in good if (g["point"].id, g["at"]) not in existing]
    by_status = {s.value: sum(1 for g in new if g["status"] == s) for s in ConditionStatus}
    days = sorted({g["at"].date() for g in new})
    return {
        "rows": new,
        "preview": {
            "readings": len(new),
            "already": len(good) - len(new),
            "points": len({g["point"].id for g in new}),
            "from": days[0].isoformat() if days else None,
            "to": days[-1].isoformat() if days else None,
            "by_status": by_status,
            "invalid": len(errors),
            "errors": errors,
        },
    }


def commit(db: Session, plant: Plant, checked: dict[str, Any], user_id: uuid.UUID) -> dict[str, int]:
    if checked["preview"]["invalid"]:
        raise ValueError("The file has errors; nothing was imported")
    now = utcnow()
    latest: dict[uuid.UUID, dict[str, Any]] = {}
    for g in checked["rows"]:
        db.execute(
            pg_insert(ConditionReading)
            .values(
                id=uuid.uuid4(),
                plant_id=plant.id,
                point_id=g["point"].id,
                value=g["value"],
                measured_at=g["at"],
                status=g["status"],
                source=SOURCE,
                recorded_by_id=user_id,
                note=g["note"],
                created_at=now,
            )
            .on_conflict_do_nothing(index_elements=["point_id", "measured_at", "source"])
        )
        cur = latest.get(g["point"].id)
        if cur is None or g["at"] > cur["at"]:
            latest[g["point"].id] = g
    moved = 0
    for g in latest.values():
        p: MeasurementPoint = g["point"]
        if p.last_reading_at is None or g["at"] > p.last_reading_at:
            p.last_value, p.last_status, p.last_reading_at = g["value"], g["status"], g["at"]
            moved += 1
    db.flush()
    audit(db, "pdm.history_import", user_id=user_id, plant_id=plant.id, data={"readings": len(checked["rows"]), "points": len(latest)})
    return {"imported": len(checked["rows"]), "points_updated": moved}
