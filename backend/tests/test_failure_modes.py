import io

from openpyxl import Workbook, load_workbook

from tests.conftest import mk_asset


def _xlsx(rows, name="codes.xlsx"):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    b = io.BytesIO()
    wb.save(b)
    return {"file": (name, b.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


HEAD = ["kind", "code", "description", "equipment_type", "category", "description_ar"]
CATALOGUE = [
    HEAD,
    ["problem", "PMP-M01", "Bearing failure", "Pumps", "mechanical", ""],
    ["problem", "PMP-M02", "Mechanical seal failure", "Pumps", "Mechanical", ""],
    ["problem", "MOT-M01", "Bearing failure", "Electric Motors", "Mechanical", ""],
    ["problem", "MOT-E01", "Winding insulation failure", "Electric Motors", "Electrical", ""],
]


def test_catalogue_import_suggests_modes_for_the_equipment_and_pareto_counts_them(world):
    admin, planner = world.api("admin"), world.api("planner")
    bad = admin.ok(
        "post",
        "/failure-codes/import",
        params={"apply": False},
        files=_xlsx(CATALOGUE + [["oops", "X", "y"], ["problem", "PMP-M01", "again"]]),
    )
    assert (bad["new"], bad["invalid"]) == (4, 2)
    assert planner.post("/failure-codes/import", params={"apply": True}, files=_xlsx(CATALOGUE)).status_code == 403
    assert admin.ok("post", "/failure-codes/import", params={"apply": True}, files=_xlsx(CATALOGUE))["new"] == 4
    upd = admin.ok(
        "post", "/failure-codes/import", params={"apply": True}, files=_xlsx([HEAD, ["problem", "PMP-M02", "Seal leak", "Pumps", "", ""]])
    )
    assert (upd["new"], upd["updated"]) == (0, 1)
    sheet = load_workbook(io.BytesIO(admin.get("/failure-codes/export.xlsx").content)).active
    assert [c.value for c in sheet[1]] == HEAD

    line = mk_asset(planner, world, "PL-1")
    pump = planner.ok(
        "post",
        "/assets",
        {
            "code": "PL-1-P01",
            "name": "Cooling water pump",
            "department_id": str(world.depts["PM"].id),
            "parent_id": line["id"],
            "level": "unit",
            "criticality": "B",
        },
        status=201,
    )
    brg = mk_asset(planner, world, "PL-1-P01-BRG", level="component", parent=pump["id"])
    preview = planner.ok("post", "/assets/equipment-types/classify", params={"apply": False})
    assert preview["by_type"].get("Pumps", 0) >= 1 and "PL-1-P01 — Cooling water pump" in preview["samples"]["Pumps"]
    planner.ok("post", "/assets/equipment-types/classify", params={"apply": True})
    assert planner.ok("get", f"/assets/{pump['id']}")["equipment_type"] == "Pumps"
    assert "Pumps" in planner.ok("get", "/assets/equipment-types")

    codes = world.api("tech").ok("get", "/failure-codes", params={"asset_id": brg["id"]})  # component inherits the pump's type
    first = [c["code"] for c in codes if c["suggested"]]
    assert first == ["PMP-M01", "PMP-M02"] and [c["code"] for c in codes[:2]] == first

    by_code = {c["code"]: c["id"] for c in codes}
    for asset, code in ((pump, "PMP-M01"), (pump, "PMP-M01"), (brg, "PMP-M02"), (line, "MOT-M01")):
        wo = planner.ok(
            "post", "/work-orders", {"type": "corrective", "priority": "P3", "asset_id": asset["id"], "title": "Fault"}, status=201
        )
        planner.ok("patch", f"/work-orders/{wo['id']}", {"problem_code_id": by_code[code], "downtime_hours": 2, "version": wo["version"]})
    q = {"start": "2000-01-01", "end": "2100-01-01"}
    modes = world.api("manager").ok("get", "/kpis/failure-modes", params=q)
    assert modes["total"] == 4
    top = modes["rows"][0]
    assert (top["failure_mode"], top["work_orders"], top["downtime_hours"]) == ("Bearing failure", 3, 6.0)  # pump + motor bearings added up
    assert top["equipment_types"] == ["Electric Motors", "Pumps"] and top["share_pct"] == 75.0
    codes_view = world.api("manager").ok("get", "/kpis/failure-modes", params={**q, "by": "code", "under": pump["id"]})
    assert [(r["code"], r["work_orders"]) for r in codes_view["rows"]] == [("PMP-M01", 2), ("PMP-M02", 1)]


def test_request_with_standard_symptom_carries_codes_to_the_work_order(world):
    admin, planner = world.api("admin"), world.api("planner")
    rows = [HEAD, ["symptom", "S-NOISE", "Abnormal noise", "", "", ""], ["problem", "GEN-BRG", "Bearing failure", "", "Mechanical", ""]]
    admin.ok("post", "/failure-codes/import", params={"apply": True}, files=_xlsx(rows))
    codes = {c["code"]: c for c in world.api("tech").ok("get", "/failure-codes")}
    assert codes["S-NOISE"]["kind"] == "symptom"
    line = mk_asset(planner, world, "SYM-1")
    tech = world.api("tech")
    bad = {"asset_id": line["id"], "title": "Abnormal noise – SYM-1", "symptom_code_id": codes["GEN-BRG"]["id"]}
    assert tech.post("/work-requests", json=bad).status_code == 422  # a failure mode is not a symptom
    req = tech.ok(
        "post", "/work-requests", {**bad, "symptom_code_id": codes["S-NOISE"]["id"], "problem_code_id": codes["GEN-BRG"]["id"]}, status=201
    )
    r = planner.ok("post", f"/work-requests/{req['id']}/triage", {"accept": True})
    wo = planner.ok("get", f"/work-orders/{r['work_order_id']}")
    assert (wo["symptom_code_id"], wo["problem_code_id"]) == (codes["S-NOISE"]["id"], codes["GEN-BRG"]["id"])
    sym = world.api("manager").ok("get", "/kpis/failure-modes", params={"start": "2000-01-01", "end": "2100-01-01", "by": "symptom"})
    assert [(x["code"], x["work_orders"]) for x in sym["rows"]] == [("S-NOISE", 1)]


def test_equipment_type_from_name():
    from app.services.failure_modes import guess_type

    cases = {
        "Cooling Pump Gearbox": "Gearboxes",
        "Transformer 7000KVA (HT Motor)": "Transformers",
        "Felt roll drive motor": "Electric Motors",
        "Press Roll No. 2": "Rolls",
        "Guide rollers": "Rolls",
        "Felt guide roll": "Rolls",
        "Dryer Cylinder 12": "Dryer Cylinders",
        "Hydraulic cylinder for nip": "Hydraulic & Pneumatic Units",
        "Pick-up felt": "Felts & Fabrics",
        "Creping doctor": "Doctors & Showers",
        "High pressure showers": "Doctors & Showers",
        "Headbox": "Headboxes",
        "Machine chest agitator": "Agitators & Mixers",
        "Machine chest": "Tanks & Chests",
        "Double disc refiner": "Refiners & Pulpers",
        "Pulper W.Water Pump": "Pumps",
        "Pressure screen": "Screens & Cleaners",
        "Overhead crane 10t": "Cranes & Hoists",
        "Wire Part (Forming Section)": None,  # a section, not a fabric
        "Rolling bearing housing": None,
    }
    assert {n: guess_type(n) for n in cases} == cases
