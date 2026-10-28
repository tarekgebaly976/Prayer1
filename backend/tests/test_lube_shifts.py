from datetime import timedelta

from app.services.common import plant_today
from tests.conftest import mk_asset


def _lube_setup(world):
    mgr, rel = world.api("manager"), world.api("reliability")
    line = mk_asset(mgr, world, "PM1")
    ep2 = rel.ok("post", "/lubrication/lubricants", {"code": "EP2", "manufacturer": "SKF", "product": "LGEP 2", "kind": "grease",
                                                      "specification": "NLGI 2"}, status=201)
    alt = rel.ok("post", "/lubrication/lubricants", {"code": "EP2B", "manufacturer": "Shell", "product": "Gadus S2", "kind": "grease"}, status=201)
    other = rel.ok("post", "/lubrication/lubricants", {"code": "PU", "manufacturer": "X", "product": "Polyurea", "kind": "grease"}, status=201)
    route = rel.ok("post", "/lubrication/routes", {"code": "R1", "name": "Wire", "department_id": str(world.depts["PM"].id), "lead_days": 1}, status=201)
    today = plant_today(world.plant)
    pts = [rel.ok("post", "/lubrication/points", {"code": f"LP{i}", "asset_id": line["id"], "lubricant_id": ep2["id"], "quantity": 40,
                                                  "unit": "g", "frequency_days": 14, "route_id": route["id"], "route_seq": i,
                                                  "next_due_date": today.isoformat()}, status=201) for i in (1, 2)]
    return line, ep2, alt, other, route, pts, today


def test_lubrication_route_record_and_next_due(world):
    """Acceptance D."""
    line, ep2, alt, other, route, pts, today = _lube_setup(world)
    rel, tech, sup = world.api("reliability"), world.api("tech"), world.api("mech_sup")
    g = rel.ok("post", f"/lubrication/routes/{route['id']}/generate", {})
    assert g["number"]
    assert rel.ok("post", f"/lubrication/routes/{route['id']}/generate", {})["work_order_id"] is None  # idempotent
    wo_id = g["work_order_id"]
    sup.ok("post", f"/work-orders/{wo_id}/assign", {"user_ids": [str(world.users["tech"].id)]})
    tech.ok("post", f"/work-orders/{wo_id}/status", {"to": "in_progress"})
    tasks = tech.get(f"/work-orders/{wo_id}/detail").json()["lube_tasks"]
    assert len(tasks) == 2
    tech.ok("post", f"/work-orders/{wo_id}/status", {"to": "awaiting_review"})
    r = sup.post(f"/work-orders/{wo_id}/status", {"to": "completed"})
    assert r.status_code == 409 and "lubrication task" in r.text
    # unapproved substitute is rejected
    r = tech.post(f"/lubrication/tasks/{tasks[0]['id']}/record", {"quantity": 40, "lubricant_used_id": other["id"]})
    assert r.status_code == 409 and "approved alternative" in r.text
    assert tech.post(f"/lubrication/lubricants/{ep2['id']}/alternatives", {"alternative_id": alt["id"], "compatibility_note": "x" * 12}).status_code == 403
    rel.ok("post", f"/lubrication/lubricants/{ep2['id']}/alternatives", {"alternative_id": alt["id"],
                                                                         "compatibility_note": "Lithium/lithium compatible per OEM letter 2026-04"}, status=201)
    res = tech.ok("post", f"/lubrication/tasks/{tasks[0]['id']}/record", {"quantity": 40, "lubricant_used_id": alt["id"]})
    assert res["next_due_date"] == (today + timedelta(days=14)).isoformat()
    res = tech.ok("post", f"/lubrication/tasks/{tasks[1]['id']}/record", {"quantity": 60})
    assert res["warnings"]  # >20% over-greasing flagged
    sup.ok("post", f"/work-orders/{wo_id}/status", {"to": "completed"})
    p = rel.get("/lubrication/points", params={"route_id": route["id"]}).json()
    assert all(x["next_due_date"] == (today + timedelta(days=14)).isoformat() and x["last_done_date"] == today.isoformat() for x in p)
    by_qr = tech.get(f"/lubrication/points/by-qr/{p[0]['qr_token']}").json()
    assert by_qr["code"] == p[0]["code"]
    k = world.api("manager").get("/kpis/lubrication_compliance", params={"start": today.isoformat(), "end": today.isoformat()}).json()
    assert k["value"] == 100.0 and k["denominator"] == 2


def test_missed_lubrication_flagged(world):
    from app.core.db import SessionLocal
    from app.models.tenancy import Plant
    from app.services.lubrication import flag_missed
    line, ep2, alt, other, route, pts, today = _lube_setup(world)
    world.api("reliability").ok("post", f"/lubrication/routes/{route['id']}/generate", {})
    with SessionLocal() as db, db.begin():
        assert flag_missed(db, db.get(Plant, world.plant.id), today + timedelta(days=5)) == 2
    with SessionLocal() as db, db.begin():
        assert flag_missed(db, db.get(Plant, world.plant.id), today + timedelta(days=6)) == 0  # notified once
    notes = world.api("reliability").get("/notifications").json()["items"]
    assert sum(1 for n in notes if n["kind"] == "lube_missed") == 2


def test_roster_pattern_and_handover(world):
    """Acceptance E."""
    mgr, tech, sup = world.api("manager"), world.api("tech"), world.api("mech_sup")
    today = plant_today(world.plant)
    pat = mgr.ok("post", "/shifts/patterns", {"code": "DDNNOO", "name": "6-day", "cycle": ["D", "D", "N", "N", "OFF", "OFF"]}, status=201)
    assert mgr.post("/shifts/patterns", {"code": "BAD", "name": "x", "cycle": ["Q"]}).status_code == 422
    team = mgr.ok("post", "/shifts/teams", {"code": "A", "name": "Crew A", "pattern_id": pat["id"], "pattern_anchor_date": today.isoformat(),
                                            "member_ids": [str(world.users["tech"].id), str(world.users["tech2"].id)],
                                            "supervisor_id": str(world.users["mech_sup"].id)}, status=201)
    # tech2 has leave on day 1: pattern generation must not overwrite it
    d1 = (today + timedelta(days=1)).isoformat()
    mgr.ok("put", "/shifts/roster", {"user_id": str(world.users["tech2"].id), "work_date": d1, "kind": "leave", "hours": 0})
    n = mgr.ok("post", f"/shifts/teams/{team['id']}/generate-roster", {"start": today.isoformat(), "end": (today + timedelta(days=5)).isoformat()})
    assert n["created"] == 7  # 4 working days × 2 people − 1 leave day
    roster = mgr.get("/shifts/roster", params={"start": today.isoformat(), "end": (today + timedelta(days=5)).isoformat()}).json()
    assert any(r["kind"] == "leave" and r["work_date"] == d1 for r in roster)
    # conflicting manual entry requires explicit replace
    r = mgr.put("/shifts/roster", {"user_id": str(world.users["tech"].id), "work_date": today.isoformat(), "shift_id": str(world.shifts["N"].id)})
    assert r.status_code == 409
    mgr.ok("put", "/shifts/roster", {"user_id": str(world.users["tech"].id), "work_date": today.isoformat(), "shift_id": str(world.shifts["D"].id),
                                     "overtime_hours": 2, "replace": True})
    mp = mgr.get("/shifts/manpower", params={"start": today.isoformat(), "end": today.isoformat()}).json()
    assert mp[0]["people"] == 2 and mp[0]["total_supply"] == 18
    # emergency open breakdown should appear in handover
    line = mk_asset(mgr, world, "PM1")
    mgr.ok("post", "/work-orders", {"asset_id": line["id"], "type": "emergency", "title": "Felt run trip"}, status=201)
    h = tech.ok("post", "/shifts/handovers", {"department_id": str(world.depts["PM"].id), "shift_date": today.isoformat(),
                                              "shift_id": str(world.shifts["D"].id), "incoming_user_id": str(world.users["tech2"].id),
                                              "summary": "Felt run trip still open", "safety_concerns": "Steam leak near dryer 12",
                                              "pending_permits": "Hot work permit HW-118"}, status=201)
    assert tech.post(f"/shifts/handovers/{h['id']}/confirm", {}).status_code == 403  # technician lacks confirm permission
    tech.ok("post", f"/shifts/handovers/{h['id']}/submit")
    detail = sup.get(f"/shifts/handovers/{h['id']}").json()
    assert len(detail["open_breakdowns_at_submit"]) == 1
    assert tech.patch(f"/shifts/handovers/{h['id']}", {"summary": "edit after submit"}).status_code == 409
    sup.ok("post", f"/shifts/handovers/{h['id']}/confirm", {"comment": "Received"})
    assert sup.get(f"/shifts/handovers/{h['id']}").json()["status"] == "confirmed"
    assert any(n["kind"] == "handover" for n in world.api("tech2").get("/notifications").json()["items"])


def test_supervisor_cannot_confirm_own_handover(world):
    sup = world.api("mech_sup")
    today = plant_today(world.plant)
    h = sup.ok("post", "/shifts/handovers", {"department_id": str(world.depts["PM"].id), "shift_date": today.isoformat(),
                                             "shift_id": str(world.shifts["N"].id), "summary": "quiet night"}, status=201)
    sup.ok("post", f"/shifts/handovers/{h['id']}/submit")
    assert sup.post(f"/shifts/handovers/{h['id']}/confirm", {}).status_code == 403


def _roster_xlsx(rows):
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return {"file": ("roster.xlsx", b.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def test_roster_import_from_plant_grid_with_shift_codes(world):
    import io

    from openpyxl import load_workbook

    admin, manager = world.api("admin"), world.api("manager")
    d = next(s for s in admin.ok("get", "/shifts/definitions") if s["code"] == "D")
    assert manager.ok("patch", f"/shifts/definitions/{d['id']}", {"name": "Day 12h", "start_time": "06:00", "end_time": "18:00", "paid_hours": 12})["paid_hours"] == 12
    manager.ok("post", "/shifts/definitions", {"code": "G", "name": "General", "kind": "general", "start_time": "07:00", "end_time": "17:00", "paid_hours": 10}, status=201)
    grid = [
        ["", "FRI", "SAT", "SUN", "MON"],
        ["Name", 30, 31, 1, 2],  # day numbers roll over into November
        ["Tech", "D", "g", "N", "L"],
        ["Planner", "G", "G", "OFF", None],
        ["Name", 30, 31, 1, 2],
        ["Fawaz  Hazazi", "N", "N", "D", "A"],
    ]
    q = {"start": "2026-10-26"}
    p = manager.ok("post", "/shifts/roster/import/preview", params=q, files=_roster_xlsx(grid))
    assert (p["people"], p["matched"], p["missing"], p["entries"], p["invalid"]) == (3, 2, ["Fawaz  Hazazi"], 11, 0)
    assert (p["from"], p["to"]) == ("2026-10-30", "2026-11-02")
    bad = manager.ok("post", "/shifts/roster/import/preview", params=q, files=_roster_xlsx(grid[:3] + [["Planner", "Z", "", "", ""]]))
    assert bad["invalid"] == 1 and "unknown code 'Z'" in bad["errors"][0]["errors"][0]

    assert manager.post("/shifts/roster/import/commit", params=q, files=_roster_xlsx(grid)).status_code == 422  # missing person
    assert manager.post("/shifts/roster/import/commit", params={**q, "create_missing": True}, files=_roster_xlsx(grid)).status_code == 403
    r = admin.ok("post", "/shifts/roster/import/commit", params={**q, "create_missing": True}, files=_roster_xlsx(grid))
    assert r == {"saved": 10, "removed": 0, "people_created": 1}

    roster = admin.ok("get", "/shifts/roster", params={"start": "2026-10-30", "end": "2026-11-02"})
    tech = {e["work_date"]: e for e in roster if e["user_id"] == str(world.users["tech"].id)}
    assert (tech["2026-10-30"]["shift"], float(tech["2026-10-30"]["hours"])) == ("D", 12.0)
    assert tech["2026-10-31"]["shift"] == "G" and tech["2026-11-02"]["kind"] == "leave"
    mp = {m["date"]: m for m in admin.ok("get", "/shifts/manpower", params={"start": "2026-10-30", "end": "2026-10-30"})}
    assert mp["2026-10-30"]["people"] == 3 and mp["2026-10-30"]["total_supply"] == 30  # D 12 + G 10 + N 8

    tpl = load_workbook(io.BytesIO(admin.get("/shifts/roster/template.xlsx", params={"start": "2026-10-30", "end": "2026-11-02"}).content)).active
    rows = {r[0]: r[1:] for r in tpl.iter_rows(min_row=3, values_only=True)}
    assert rows["Tech"] == ("D", "G", "N", "L") and rows["Fawaz  Hazazi"][3] == "A"  # template comes back pre-filled

    off = [["Name", 30], ["Tech", "OFF"]]
    assert admin.ok("post", "/shifts/roster/import/commit", params=q, files=_roster_xlsx(off)) == {"saved": 0, "removed": 1, "people_created": 0}


def test_manpower_counts_unscheduled_work_on_its_due_date(world):
    planner = world.api("planner")
    line = mk_asset(planner, world, "MP-LINE")
    day = plant_today(world.plant) + timedelta(days=3)
    planner.ok("post", "/work-orders", {"asset_id": line["id"], "title": "Gearbox inspection", "type": "corrective",
                                        "due_date": day.isoformat(), "estimated_hours": 6}, status=201)
    mp = planner.ok("get", "/shifts/manpower", params={"start": day.isoformat(), "end": day.isoformat()})[0]
    assert mp["total_demand"] == 6 and mp["total_supply"] == 0 and mp["flags"]  # no one rostered for 6 h of work
