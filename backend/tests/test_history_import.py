import io

from openpyxl import Workbook, load_workbook

from tests.conftest import mk_asset

WO_HEADER = [
    "external_ref",
    "asset_code",
    "type",
    "title",
    "description",
    "trade",
    "failure_mode",
    "technician",
    "priority",
    "original_status",
    "breakdown",
    "reported_at",
    "started_at",
    "completed_at",
    "closed_at",
    "downtime_hours",
    "repair_minutes",
    "findings",
]
PART_HEADER = ["wo_external_ref", "part_code", "description", "quantity", "uom", "unit_cost", "used_on"]


def _xlsx(rows, name="f.xlsx"):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return {"file": (name, b.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def test_history_work_orders_and_parts_feed_reports_and_kpis(world):
    admin, planner, store = world.api("admin"), world.api("planner"), world.api("store")
    line = mk_asset(planner, world, "SP", dept="SP")
    pump = mk_asset(planner, world, "SP-ONP", level="unit", parent=line["id"], dept="SP")
    store.ok("post", "/inventory/parts", {"code": "4000038", "description": "Tape electrical", "uom": "PCS", "unit_price": 30}, status=201)
    rows = [
        WO_HEADER,
        [
            "RM-38692",
            "SP-ONP",
            "",
            "Gland leakage",
            "Pump gland",
            "Mechanical",
            "UN-BALANCED",
            "Zande M.",
            "High",
            "Complete",
            "yes",
            "2026-01-04 08:07",
            "2026-01-04 08:30",
            "2026-01-04 16:30",
            "",
            "",
            480,
            "DONE",
        ],
        [
            "SO-240001258",
            "SP-ONP",
            "shutdown",
            "Planned shutdown",
            "",
            "Electrical",
            "",
            "V. Hernandez",
            "",
            "closed",
            "no",
            "2026-05-19 07:00",
            "",
            "",
            "",
            24,
            "",
            "",
        ],
        [
            "RM-38714",
            "SP-ONP",
            "workshop",
            "Railing support repair",
            "",
            "Workshop",
            "RE-GRINDING",
            "",
            "",
            "Open",
            "false",
            "2026-01-05 09:28",
            "",
            "",
            "",
            "",
            "",
            "",
        ],
    ]
    bad = rows + [["RM-38692", "NOPE", "teleport", "", "", "", "", "", "Whenever", "", "", "05/01/2026", "", "", "", "", "", ""]]
    assert planner.post("/imports/wo_history/preview", files=_xlsx(bad)).status_code == 403  # administrators only
    p = admin.ok("post", "/imports/wo_history/preview", files=_xlsx(bad))
    assert p["invalid"] == 1 and p["plans"] == 1  # plans = breakdowns
    msg = " ".join(p["errors"][0]["errors"])
    assert all(x in msg for x in ("duplicate external_ref", "not found", "unknown type", "title", "priority", "reported_at"))
    assert admin.ok("post", "/imports/wo_history/commit", files=_xlsx(rows))["created"] == 3

    wos = {w["external_ref"]: w for w in admin.get("/work-orders", params={"limit": 50, "status": "closed"}).json()["items"]}
    bd = wos["RM-38692"]
    assert bd["number"] == "MIG-RM-38692" and bd["source"] == "migrated" and bd["type"] == "emergency" and bd["priority"] == "P2"
    assert bd["status"] == "closed" and float(bd["downtime_hours"]) == 8.0 and float(bd["actual_labor_hours"]) == 8.0
    assert wos["RM-38714"]["type"] == "corrective"  # workshop jobs count as maintenance work
    assert admin.patch(f"/work-orders/{bd['id']}", {"title": "x"}).status_code == 409  # history is read-only
    again = admin.ok("post", "/imports/wo_history/preview", files=_xlsx(rows[:2]))
    assert "already imported" in again["errors"][0]["errors"][0]

    parts = [
        PART_HEADER,
        ["SO-240001258", "04000038", "TAPE ELECTRICAL", 3, "PCS", 29.86, "2026-05-19"],
        ["SO-240001258", "X0402063", "WORM GEAR FOR LASER", 1, "", 1200, "2026-05-19"],
    ]
    pp = admin.ok("post", "/imports/parts_history/preview", files=_xlsx(parts + [["SO-999", "", "", 0, "", "", ""]]))
    assert pp["invalid"] == 1 and pp["plans"] == 1  # one line matched an item in the spare-parts list (leading zero ignored)
    assert admin.ok("post", "/imports/parts_history/commit", files=_xlsx(parts))["created"] == 2
    rep = admin.ok("get", "/reports/parts-used", params={"asset_id": pump["id"]})
    assert rep["total"] == 2 and rep["total_cost"] == 1289.58
    assert {r["description"] for r in rep["items"]} == {"Tape electrical", "X0402063 - WORM GEAR FOR LASER"}

    k = admin.ok("get", "/kpis/mttr", params={"start": "2026-01-01", "end": "2026-01-31"})
    assert k["available"] and k["value"] == 8.0  # the migrated breakdown is a failure event
    c = admin.ok("get", "/kpis/cost_by_department", params={"start": "2026-05-01", "end": "2026-05-31"})
    assert c["value"] == 1289.58 and c["details"]["by_department"] == {"SP": 1289.58}  # parts used count as material cost


def test_company_kpi_scorecard_template_roundtrip(world):
    admin, viewer = world.api("admin"), world.api("viewer")
    tpl = viewer.get("/kpis/scorecard/template.xlsx", params={"year": 2026})
    wb = load_workbook(io.BytesIO(tpl.content))
    assert wb.sheetnames == ["KPI 2026", "instructions", "codes"]
    head = [c.value for c in wb.active[1]]
    assert head[:8] == ["kpi_name", "department", "unit", "higher_is_better", "cmms_kpi", "frequency", "row", "2026-01"]
    rows = [
        head,
        ["Major Breakdown", "", "count", "no", "", "Monthly", "Target", 2, 2, 2],
        ["Major Breakdown", "", "count", "no", "", "Monthly", "Actual", 0, 2, 3],
        ["MTBF - PM", "PM", "h", "yes", "mtbf", "Monthly", "Target", 35, 35],
        ["MTBF - PM", "PM", "h", "yes", "mtbf", "Monthly", "Actual", 110, "x"],
    ]
    bad = rows + [["Oops", "ZZ", "", "maybe", "nope", "", "Forecast", "abc"]]
    assert viewer.post("/kpis/scorecard/preview", files=_xlsx(bad)).status_code == 403
    p = admin.ok("post", "/kpis/scorecard/preview", files=_xlsx(bad))
    assert (p["kpis"], p["values"], p["cleared"], p["invalid"]) == (3, 9, 1, 1)
    msg = " ".join(p["errors"][0]["errors"])
    assert all(x in msg for x in ("department ZZ", "higher_is_better", "cmms_kpi nope", "Target or Actual"))
    assert admin.post("/kpis/scorecard/commit", files=_xlsx(bad)).status_code == 422
    assert admin.ok("post", "/kpis/scorecard/commit", files=_xlsx(rows)) == {"kpis_created": 2, "saved": 9, "cleared": 1}

    cards = {c["name"]: c for c in viewer.ok("get", "/kpis/scorecard", params={"year": 2026})}
    mb, mtbf = cards["Major Breakdown"], cards["MTBF - PM"]
    assert mb["department"] is None and mb["higher_is_better"] is False and mb["months"][2] == {"target": 2.0, "actual": 3.0}
    assert mtbf["department"] == "PM" and mtbf["kpi_code"] == "mtbf" and mtbf["months"][0] == {"target": 35.0, "actual": 110.0}
    trend = viewer.ok("get", "/kpis/trend/mtbf", params={"months": 24, "department_id": str(world.depts["PM"].id)})
    jan = next(t for t in trend if t["period"] == "2026-01")
    assert (jan["manual"], jan["target"]) == (110.0, 35.0)  # linked company KPI shown next to the calculated value
    assert all(t["manual"] is None for t in viewer.ok("get", "/kpis/trend/mtbf", params={"months": 24}))  # plant level: not linked

    again = load_workbook(io.BytesIO(viewer.get("/kpis/scorecard/template.xlsx", params={"year": 2026}).content)).active
    filled = [list(r) for r in again.iter_rows(values_only=True)]
    assert filled[3][:9] == ["MTBF - PM", "PM", "h", "yes", "mtbf", "Monthly", "Target", 35, 35]  # template comes back pre-filled
    assert viewer.delete(f"/kpis/scorecard/{mb['id']}").status_code == 403
    assert admin.delete(f"/kpis/scorecard/{mb['id']}").status_code == 204
    assert [c["name"] for c in viewer.ok("get", "/kpis/scorecard", params={"year": 2026})] == ["MTBF - PM"]


def test_nightly_kpi_snapshot_job_saves_values(world):
    """The worker's kpi_snapshot job writes rows that reference the KPI catalogue (seeded by sync_rbac)."""
    from sqlalchemy import func, select

    from app.api.routers.kpi import snapshot_month
    from app.core.db import SessionLocal
    from app.models.system import KPISnapshot
    from app.models.tenancy import Plant

    with SessionLocal() as db, db.begin():
        plant = db.get(Plant, world.plant.id)
        assert snapshot_month(db, plant, 2026, 1) == 20
    with SessionLocal() as db:
        assert db.execute(select(func.count()).select_from(KPISnapshot).where(KPISnapshot.scope_key == "plant")).scalar() == 20


def test_company_kpi_month_entry_in_app(world):
    admin, viewer = world.api("admin"), world.api("viewer")
    head = ["kpi_name", "department", "unit", "higher_is_better", "cmms_kpi", "frequency", "row", "2026-01"]
    admin.ok(
        "post",
        "/kpis/scorecard/commit",
        files=_xlsx(
            [head, ["MTTR - PM", "PM", "h", "no", "mttr", "", "Target", 2], ["Breakdowns", "", "count", "no", "", "", "Actual", 1]]
        ),
    )
    cards = {c["name"]: c["id"] for c in viewer.ok("get", "/kpis/scorecard", params={"year": 2026})}
    body = {
        "period": "2026-11",
        "values": [{"id": cards["MTTR - PM"], "actual": 1.8, "target": 2}, {"id": cards["Breakdowns"], "actual": 0, "target": None}],
    }
    assert viewer.put("/kpis/scorecard/month", body).status_code == 403
    assert admin.ok("put", "/kpis/scorecard/month", body) == {"changed": 2}
    assert admin.ok("put", "/kpis/scorecard/month", body) == {"changed": 0}  # same values again
    got = {c["name"]: c["months"] for c in viewer.ok("get", "/kpis/scorecard", params={"year": 2026})}
    assert got["MTTR - PM"][10] == {"target": 2.0, "actual": 1.8} and got["Breakdowns"][10] == {"target": None, "actual": 0.0}
    assert got["MTTR - PM"][0] == {"target": 2.0, "actual": None}  # other months untouched
    clear = {"period": "2026-11", "values": [{"id": cards["Breakdowns"], "actual": None, "target": None}]}
    assert admin.ok("put", "/kpis/scorecard/month", clear) == {"changed": 1}
    assert viewer.ok("get", "/kpis/scorecard", params={"year": 2026})[1]["months"][10] == {"target": None, "actual": None}
    assert admin.put("/kpis/scorecard/month", {"period": "2026-13", "values": []}).status_code == 422
    bad = {"period": "2026-11", "values": [{"id": "00000000-0000-0000-0000-000000000000", "actual": 1}]}
    assert admin.put("/kpis/scorecard/month", bad).status_code == 422
