import io

from openpyxl import Workbook, load_workbook

from tests.conftest import login

COLS = ["full_name", "email", "role", "department_code", "trade", "employee_no", "hourly_rate", "reset_password", "old_email"]


def _xlsx(rows):
    wb = Workbook()
    for r in [COLS, *rows]:
        wb.active.append(r)
    b = io.BytesIO()
    wb.save(b)
    return {"file": ("users.xlsx", b.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def test_users_from_excel_add_update_and_hand_out_passwords(world):
    admin = world.api("admin")
    sheet = load_workbook(io.BytesIO(admin.get("/users/import/template.xlsx").content))["users"]
    assert [c.value for c in sheet[1]] == COLS
    listed = {r[1]: r for r in sheet.iter_rows(min_row=2, values_only=True)}
    tech = world.users["tech"]
    assert listed[tech.email][8] == tech.email  # existing users come pre-filled with old_email

    rows = [
        ["Rahim Uddin", "rahim@example.com", "technician", "SP", "mechanical", "E100", 25, "", ""],
        ["Ana Cruz", "ana@example.com", "Mechanical Supervisor", "", "", "", "", "", ""],
        [tech.full_name, "tech.new@example.com", "technician", "PM", "electrical", "", "", "", tech.email],
    ]
    bad = admin.ok(
        "post",
        "/users/import",
        params={"apply": False},
        files=_xlsx(
            rows
            + [
                ["X", "not-an-email", "technician", "", "", "", "", "", ""],
                ["Bad Role", "b@example.com", "pilot", "ZZ", "", "", "-1", "", ""],
                ["Boss", "boss@example.com", "system_admin", "", "", "", "", "", ""],
            ]
        ),
    )
    assert (bad["new"], bad["updated"], bad["invalid"]) == (3, 1, 2)  # the system administrator may give the system_admin role
    assert world.api("manager").post("/users/import", params={"apply": True}, files=_xlsx(rows)).status_code == 403

    done = admin.ok("post", "/users/import", params={"apply": True}, files=_xlsx(rows))
    assert (done["created"], done["changed"]) == (2, 1)
    creds = {c["email"]: c["temporary_password"] for c in done["credentials"]}
    assert set(creds) == {"rahim@example.com", "ana@example.com"}
    rahim = login("rahim@example.com", world.plant.id, password=creds["rahim@example.com"])
    me = rahim.get("/auth/me").json()
    assert me["full_name"] == "Rahim Uddin"
    members = {m["user"]["email"]: m for m in admin.ok("get", "/users")}
    assert members["rahim@example.com"]["role_code"] == "technician" and members["rahim@example.com"]["department_id"] == str(
        world.depts["SP"].id
    )
    assert members["tech.new@example.com"]["user"]["trade"] == "electrical" and tech.email not in members  # e-mail changed in place

    again = admin.ok(
        "post",
        "/users/import",
        params={"apply": False},
        files=_xlsx([["Rahim Uddin", "rahim@example.com", "technician", "SP", "mechanical", "E100", 25, "", "rahim@example.com"]]),
    )
    assert (again["new"], again["updated"], again["unchanged"]) == (0, 0, 1)
    reset = admin.ok(
        "post",
        "/users/import",
        params={"apply": True},
        files=_xlsx([["Rahim Uddin", "rahim@example.com", "technician", "SP", "mechanical", "E100", 25, "yes", "rahim@example.com"]]),
    )
    assert len(reset["credentials"]) == 1 and reset["credentials"][0]["temporary_password"] != creds["rahim@example.com"]
    assert rahim.get("/auth/me").status_code == 401  # old sessions signed out after a reset
