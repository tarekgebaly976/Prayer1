import io

from openpyxl import load_workbook

from tests.conftest import mk_asset


def _wo(planner, asset_id, title="Replace bearing"):
    return planner.ok("post", "/work-orders", {"asset_id": asset_id, "type": "corrective", "title": title}, status=201)


def test_parts_used_without_stock_cost_and_report_by_machine(world):
    planner, tech, store, prod = world.api("planner"), world.api("tech"), world.api("store"), world.api("production")
    line = mk_asset(planner, world, "PM1")
    dryer = mk_asset(planner, world, "PM1-DRY", level="unit", parent=line["id"])
    press = mk_asset(planner, world, "PM1-PRS", level="unit", parent=line["id"])
    part = store.ok(
        "post", "/inventory/parts", {"code": "BRG-6315", "description": "Bearing 6315 C3", "uom": "PCS", "unit_price": 250}, status=201
    )
    wo1, wo2 = _wo(planner, dryer["id"]), _wo(planner, press["id"], "Seal leak")

    # No stock exists anywhere: recording usage must still work and must not touch stock levels.
    r = tech.ok("post", f"/work-orders/{wo1['id']}/parts-used", {"part_id": part["id"], "quantity": 2, "note": "DE + NDE"}, status=201)
    assert r["asset_code"] == "PM1-DRY" and float(r["total_cost"]) == 500 and r["wo_material_cost"] == 500
    tech.ok(
        "post",
        f"/work-orders/{wo2['id']}/parts-used",
        {"description": "Mechanical seal 45mm (local purchase)", "quantity": 1, "unit_cost": 1200, "uom": "SET"},
        status=201,
    )
    assert tech.post(f"/work-orders/{wo1['id']}/parts-used", {"quantity": 1}).status_code == 422  # part or text needed
    assert prod.post(f"/work-orders/{wo1['id']}/parts-used", {"part_id": part["id"], "quantity": 1}).status_code == 403
    detail = tech.ok("get", f"/work-orders/{wo1['id']}/detail")
    assert [(p["part_code"], float(p["quantity"])) for p in detail["parts_used"]] == [("BRG-6315", 2.0)]
    stock = store.ok("get", f"/inventory/parts/{part['id']}")
    assert float(stock.get("qty_on_hand", 0) or 0) == 0

    rep = planner.ok("get", "/reports/parts-used", params={"asset_id": line["id"]})  # whole line = both machines
    assert rep["total"] == 2 and rep["total_cost"] == 1700
    only_dryer = planner.ok("get", "/reports/parts-used", params={"asset_id": dryer["id"]})
    assert only_dryer["total"] == 1 and only_dryer["items"][0]["work_order"] == wo1["number"]
    by_part = planner.ok("get", "/reports/parts-used", params={"part_id": part["id"]})
    assert by_part["by_part"][0]["quantity"] == 2

    x = planner.get("/exports/parts-used.xlsx")
    ws = load_workbook(io.BytesIO(x.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][1] == "Work order" and len(rows) == 3 and {r[4] for r in rows[1:]} == {"PM1-DRY", "PM1-PRS"}

    used_id = detail["parts_used"][0]["id"]
    tech.ok("delete", f"/work-orders/{wo1['id']}/parts-used/{used_id}", status=204)
    assert float(tech.ok("get", f"/work-orders/{wo1['id']}")["actual_material_cost"]) == 0


def test_parts_used_blocked_on_closed_or_cancelled_wo_and_other_plant(world):
    planner, tech = world.api("planner"), world.api("tech")
    line = mk_asset(planner, world, "PM1")
    wo = _wo(planner, line["id"])
    planner.ok("post", f"/work-orders/{wo['id']}/status", {"to": "cancelled", "comment": "dup"})
    r = tech.post(f"/work-orders/{wo['id']}/parts-used", {"description": "x", "quantity": 1})
    assert r.status_code == 409
    other = world.api("planner_b", plant=world.plant_b)
    assert other.post(f"/work-orders/{wo['id']}/parts-used", {"description": "x", "quantity": 1}).status_code in (403, 404)
